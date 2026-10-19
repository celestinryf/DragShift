"""Office, ebook and text formats (dragshift/documents.py, sheets.py, slides.py). Every fixture is
built in a temp folder: workbooks with openpyxl/odfpy, decks with python-pptx, EPUBs and Pages
files as hand-made zips (not with the code under test), PDFs with PyMuPDF. Results are checked by
reading them back with the libraries directly: cell values, sheet and slide counts, text, pages."""
import csv
import datetime as dt
import hashlib
import io
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image

from dragshift import convert, documents, report, sheets, slides
from dragshift.convert import fitz
from dragshift.util import UserError, category

REAL_IWORK = documents.iwork_to_pdf


@pytest.fixture(autouse=True)
def no_iwork(monkeypatch):
    """Tests use the built-in renderers, never a real Keynote or Pages."""
    monkeypatch.setattr(documents, "iwork_to_pdf", lambda *a, **k: False)


# ---------------------------------------------------------------- helpers
def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_text(path: Path) -> str:
    with fitz.open(path) as d:
        return "\n".join(p.get_text() for p in d)


def pdf_pages(path: Path) -> int:
    with fitz.open(path) as d:
        return d.page_count


def read_csv(path: Path, delim=",") -> list[list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f, delimiter=delim))


def leftovers(folder: Path) -> list[str]:
    return [p.name for p in folder.iterdir() if p.name.startswith(".~")]


def png_bytes(size=(200, 120), color=(30, 160, 90)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def make_workbook(path: Path) -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Name", "Qty", "Price", "Date", "Paid"])
    ws.append(["Zürich café", 3, 1.5, dt.datetime(2024, 3, 1), True])
    ws.append(["日本語, \"quoted\"", 10, 2.25, dt.datetime(2024, 3, 2, 14, 30), False])
    ws.append(["line1\nline2", None, -4, None, None])
    other = wb.create_sheet("Ünïcode ✓")
    other.append(["ключ", "значение"])
    other.append(["a", 1])
    wb.create_sheet("Empty")
    wb.save(path)
    return path


def make_deck(path: Path, hidden=True) -> Path:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly Review"
    s.placeholders[1].text = "Grüße aus Zürich 日本"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    tf = s.placeholders[1].text_frame
    tf.text = "First point"
    p = tf.add_paragraph()
    p.text, p.level = "Sub point", 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1), Inches(3))
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1), Inches(3), Inches(2)).text = "Oval"
    table = s.shapes.add_table(2, 2, Inches(1), Inches(5.5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Header cell"
    table.cell(1, 1).text = "Body cell"
    if hidden:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Secret slide"
        s._element.set("show", "0")
    prs.save(path)
    return path


def make_text_pdf(path: Path, pages=2) -> Path:
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello page {i + 1}", fontsize=24)
        page.insert_text((72, 120), "Body text line for extraction.", fontsize=11)
    doc.save(path)
    return path


def make_table_pdf(path: Path) -> Path:
    """A ruled 3x3 table PyMuPDF's find_tables can detect."""
    doc = fitz.open()
    page = doc.new_page()
    x0, y0, cw, rh = 72, 100, 120, 24
    rows = [["Item", "Count", "Cost"], ["Apples", "12", "3.50"], ["Pears", "7", "2.25"]]
    for r in range(len(rows) + 1):
        page.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(rows) * rh))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            page.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), text, fontsize=11)
    doc.save(path)
    return path


def make_epub(path: Path, chapters, title="Test Book", images=None, encryption=None, spine=True,
              repeat=1) -> Path:
    """A minimal EPUB 3 written by hand (independent of documents.write_epub). `repeat` lists the
    spine that many times over, like a hostile book that re-reads one chapter."""
    images = images or {}
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OPS/book.opf" media-type="application/oebps-package+xml"/>'
                   "</rootfiles></container>")
        if encryption:
            z.writestr("META-INF/encryption.xml", encryption)
        items, refs = [], []
        for i, body in enumerate(chapters, 1):
            z.writestr(f"OPS/text/ch{i}.xhtml",
                       '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
                       f"<head><title>c{i}</title></head><body>{body}</body></html>")
            items.append(f'<item id="ch{i}" href="text/ch{i}.xhtml" media-type="application/xhtml+xml"/>')
            refs.append(f'<itemref idref="ch{i}"/>')
        for name, data in images.items():
            z.writestr(f"OPS/{name}", data)
            items.append(f'<item id="{name.replace("/", "_")}" href="{name}" media-type="image/png"/>')
        z.writestr("OPS/book.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                   f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{title}</dc:title></metadata>'
                   f"<manifest>{''.join(items)}</manifest>"
                   f"<spine>{''.join(refs) * repeat if spine else ''}</spine></package>")
    return path


MARKDOWN = """# Report title

Some **bold**, *italic* and a [link](https://example.com/page).

- first
- second
  - nested item

1. one
2. two

```python
print("hi")
```

| Name | Qty |
|------|-----|
| Zürich | 3 |

> quoted words

<script>alert("x")</script>

[bad](javascript:alert(1))

# Second chapter

Closing text with ünïcödé.
"""


@pytest.fixture
def markdown_file(tmp_path) -> Path:
    p = tmp_path / "notes.md"
    p.write_text(MARKDOWN, encoding="utf-8")
    return p


@pytest.fixture
def docx_file(tmp_path) -> Path:
    from docx import Document
    from docx.shared import Inches
    d = Document()
    d.add_heading("Word heading", level=1)
    p = d.add_paragraph("Plain and ")
    p.add_run("bold words").bold = True
    documents._add_hyperlink(p, " a link", "https://example.org/")
    d.add_paragraph("Bullet one", style="List Bullet")
    d.add_paragraph("Bullet nested", style="List Bullet 2")
    d.add_paragraph("Number one", style="List Number")
    d.add_picture(io.BytesIO(png_bytes((300, 200))), width=Inches(2))
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 1).text = "T00", "T11"
    path = tmp_path / "word doc.docx"
    d.save(path)
    return path


RTF = (r"{\rtf1\ansi\ansicpg1252{\fonttbl\f0\fswiss Helvetica;}\f0\fs24 "
       r"Gr\'fc\'dfe aus {\b Z\'fcrich}\par Second line here\par}")


# ---------------------------------------------------------------- categories and targets
def test_categories():
    for ext in ("xlsx", "XLSX", "xlsm", "xls", "ods", "csv", "tsv"):
        assert category(f"x.{ext}") == "sheet", ext
    assert category("x.pptx") == category("x.PPT") == "slides"
    assert category("x.epub") == "epub"
    assert category("x.md") == category("x.markdown") == "markdown"
    assert category("x.rtf") == category("x.rtfd") == "rtf"
    assert category("x.pages") == "pages"
    for ext in ("txt", "html", "htm", "json", "log", "xml"):
        assert category(f"x.{ext}") == "text", ext


def test_targets():
    assert convert.targets_for("sheet", ".xlsx") == ["ODS", "CSV", "TSV", "PDF", "DOCX"]
    assert convert.targets_for("sheet", ".csv") == ["XLSX", "ODS", "TSV", "PDF", "DOCX"]
    assert convert.targets_for("sheet", ".xls") == documents.SHEET_TARGETS
    assert convert.targets_for("slides", ".pptx") == ["PDF", "JPG", "PNG"]
    for t in ("XLSX", "PPTX", "EPUB", "MD"):
        assert t in convert.targets_for("pdf", ".pdf")
    assert convert.targets_for("docx", ".docx") == ["PDF", "TXT", "MD", "HTML", "EPUB"]
    assert convert.targets_for("text", ".txt") == ["PDF", "DOCX", "HTML", "EPUB"]
    assert convert.targets_for("text", ".json") == ["PDF", "DOCX"]
    assert "MD" in convert.targets_for("text", ".html")
    assert convert.targets_for("epub", ".epub") == ["PDF", "DOCX", "TXT", "MD"]
    assert convert.targets_for("pages", ".pages") == ["PDF"]
    with pytest.raises(ValueError):
        documents.convert("x.md", "markdown", "XLSX")


def test_every_new_target_has_a_menu_hint():
    from dragshift import app
    for cat, ext in (("sheet", ".xls"), ("slides", ".pptx"), ("pdf", ".pdf"), ("docx", ".docx"),
                     ("epub", ".epub"), ("markdown", ".md"), ("rtf", ".rtf"), ("pages", ".pages"),
                     ("text", ".html"), ("text", ".txt")):
        hints = {**app.FORMAT_HINTS, **app.CATEGORY_FORMAT_HINTS.get(cat, {})}
        for t in convert.targets_for(cat, ext):
            assert hints.get(t), (cat, t)


# ---------------------------------------------------------------- spreadsheets
def test_xlsx_to_csv_one_file_per_sheet(tmp_path):
    src = make_workbook(tmp_path / "Book.xlsx")
    before = digest(src)
    outs = convert.convert(src, "sheet", "CSV")
    assert [o.name for o in outs] == ["Book - Sales.csv", "Book - Ünïcode ✓.csv"]  # empty sheet skipped
    rows = read_csv(outs[0])
    assert rows[0] == ["Name", "Qty", "Price", "Date", "Paid"]
    assert rows[1] == ["Zürich café", "3", "1.5", "2024-03-01", "TRUE"]
    assert rows[2] == ["日本語, \"quoted\"", "10", "2.25", "2024-03-02 14:30:00", "FALSE"]
    assert rows[3] == ["line1\nline2", "", "-4"]
    assert read_csv(outs[1]) == [["ключ", "значение"], ["a", "1"]]
    assert outs[0].read_bytes().startswith(b"\xef\xbb\xbf")  # BOM so Excel reads UTF-8
    assert digest(src) == before and not leftovers(tmp_path)


def test_xlsx_to_tsv_and_single_sheet_name(tmp_path):
    import openpyxl
    src = make_workbook(tmp_path / "Book.xlsx")
    outs = convert.convert(src, "sheet", "TSV")
    assert len(outs) == 2 and read_csv(outs[1], "\t")[1] == ["a", "1"]
    wb = openpyxl.Workbook()
    wb.active.append(["only", 1])
    wb.save(tmp_path / "Solo.xlsx")
    assert [o.name for o in convert.convert(tmp_path / "Solo.xlsx", "sheet", "CSV")] == ["Solo.csv"]


def test_xlsx_to_ods_round_trip(tmp_path):
    src = make_workbook(tmp_path / "Book.xlsx")
    ods = convert.convert(src, "sheet", "ODS")[0]
    with zipfile.ZipFile(ods) as z:
        assert z.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
        content = z.read("content.xml").decode("utf-8")
    assert 'table:name="Ünïcode ✓"' in content and "Zürich café" in content
    book = sheets.read_sheets(ods)
    assert [s.name for s in book] == ["Sales", "Ünïcode ✓", "Empty"]
    sales = book[0].rows
    assert sales[1] == ["Zürich café", 3, 1.5, dt.datetime(2024, 3, 1), True]
    assert sales[3] == ["line1\nline2", None, -4]
    back = convert.convert(ods, "sheet", "XLSX")[0]
    import openpyxl
    wb = openpyxl.load_workbook(back)
    assert wb.sheetnames == ["Sales", "Ünïcode ✓", "Empty"]
    assert wb["Sales"]["B3"].value == 10 and wb["Sales"]["D2"].value == dt.datetime(2024, 3, 1)
    assert wb["Sales"]["E2"].value is True


def test_csv_to_xlsx_recognises_numbers_but_keeps_codes(tmp_path):
    import openpyxl
    src = tmp_path / "data.csv"
    src.write_text('id,amount,code,note\n1,1.50,007,"a, b"\n2,-3e2,,"multi\nline"\n', encoding="utf-8")
    out = convert.convert(src, "sheet", "XLSX")[0]
    ws = openpyxl.load_workbook(out).active
    assert ws.title == "data"
    assert [c.value for c in ws[2]] == [1, 1.5, "007", "a, b"]
    assert ws["B3"].value == -300.0 and ws["C3"].value is None and ws["D3"].value == "multi\nline"
    tsv = convert.convert(src, "sheet", "TSV")[0]
    assert read_csv(tsv, "\t")[1] == ["1", "1.50", "007", "a, b"]  # CSV -> TSV keeps text exactly


def test_text_starting_with_equals_never_becomes_a_formula(tmp_path):
    """A CSV cell like =HYPERLINK(...) is text, not something Excel should run."""
    import openpyxl
    src = tmp_path / "links.csv"
    src.write_text('name,value\nx,"=HYPERLINK(""http://example.com"",""click"")"\ny,=1+1\n', encoding="utf-8")
    out = convert.convert(src, "sheet", "XLSX")[0]
    with zipfile.ZipFile(out) as z:
        assert b"<f>" not in z.read("xl/worksheets/sheet1.xml")
    ws = openpyxl.load_workbook(out).active
    assert ws["B2"].data_type == "s" and ws["B2"].value == '=HYPERLINK("http://example.com","click")'
    assert ws["B3"].data_type == "s" and ws["B3"].value == "=1+1"
    # the same for a workbook source: a text cell there stays text in the copy
    book = [sheets.Sheet("S", [["=SUM(A1:A2)", 1]])]
    again = tmp_path / "again.xlsx"
    sheets.write_xlsx(book, again)
    cell = openpyxl.load_workbook(again).active["A1"]
    assert cell.data_type == "s" and cell.value == "=SUM(A1:A2)"


def test_semicolon_csv_and_tsv_input(tmp_path):
    import openpyxl
    src = tmp_path / "europe.csv"
    src.write_text("name;price\nKäse;3,50\nBrot;2\n", encoding="cp1252")
    ws = openpyxl.load_workbook(convert.convert(src, "sheet", "XLSX")[0]).active
    assert [c.value for c in ws[2]] == ["Käse", "3,50"] and ws["B3"].value == 2
    tsv = tmp_path / "t.tsv"
    tsv.write_text("a\tb\n1\t2\n", encoding="utf-8")
    assert read_csv(convert.convert(tsv, "sheet", "CSV")[0]) == [["a", "b"], ["1", "2"]]


def test_xls_input(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Old")
    ws.write(0, 0, "Grüße")
    ws.write(0, 1, 42)
    ws.write(1, 0, 2.5)
    ws.write(1, 1, dt.datetime(2020, 1, 2), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
    ws.write(2, 0, True)
    wb.add_sheet("Second").write(0, 0, "x")
    src = tmp_path / "legacy.xls"
    wb.save(str(src))
    book = sheets.read_sheets(src)
    assert [s.name for s in book] == ["Old", "Second"]
    assert book[0].rows == [["Grüße", 42], [2.5, dt.datetime(2020, 1, 2)], [True]]
    import openpyxl
    out = convert.convert(src, "sheet", "XLSX")[0]
    assert openpyxl.load_workbook(out)["Old"]["B1"].value == 42


def test_content_beats_extension(tmp_path):
    real = make_workbook(tmp_path / "Book.xlsx")
    renamed = tmp_path / "renamed.xls"
    renamed.write_bytes(real.read_bytes())
    assert sheets.read_sheets(renamed)[0].name == "Sales"
    fake = tmp_path / "export.xls"
    fake.write_text("<html><table><tr><td>1</td></tr></table></html>")
    with pytest.raises(UserError, match="isn't a spreadsheet"):
        convert.convert(fake, "sheet", "CSV")


def test_sheet_error_paths(tmp_path, monkeypatch):
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04" + b"garbage" * 50)
    with pytest.raises(UserError, match="damaged"):
        convert.convert(broken, "sheet", "CSV")
    locked = tmp_path / "locked.xlsx"
    locked.write_bytes(sheets.OLE_MAGIC + bytes(500) + "EncryptionInfo".encode("utf-16-le") + bytes(500))
    with pytest.raises(UserError, match="password"):
        convert.convert(locked, "sheet", "CSV")
    import openpyxl
    openpyxl.Workbook().save(tmp_path / "blank.xlsx")
    with pytest.raises(UserError, match="empty"):
        convert.convert(tmp_path / "blank.xlsx", "sheet", "CSV")
    (tmp_path / "nothing.csv").write_text("")
    with pytest.raises(UserError, match="empty"):
        convert.convert(tmp_path / "nothing.csv", "sheet", "XLSX")
    big = tmp_path / "big.csv"
    big.write_text("\n".join(",".join(["1"] * 5) for _ in range(20)))
    monkeypatch.setattr(sheets, "MAX_CELLS", 50)
    with pytest.raises(UserError, match="too big"):
        convert.convert(big, "sheet", "XLSX")
    monkeypatch.setattr(sheets, "MAX_CELLS", 10_000)
    monkeypatch.setattr(sheets, "MAX_TABLE_ROWS", 5)
    for target in ("PDF", "DOCX"):
        with pytest.raises(UserError, match="CSV or XLSX"):
            convert.convert(big, "sheet", target)
    assert not list(tmp_path.glob("big*.pdf")) and not leftovers(tmp_path)


def test_encrypted_ods(tmp_path):
    src = tmp_path / "secret.ods"
    with zipfile.ZipFile(src, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("META-INF/manifest.xml", '<manifest:manifest><manifest:file-entry>'
                   '<manifest:encryption-data/></manifest:file-entry></manifest:manifest>')
        z.writestr("content.xml", b"\x00\x01encrypted")
    with pytest.raises(UserError, match="password"):
        convert.convert(src, "sheet", "CSV")


def test_ods_repeated_cells_dont_blow_up(tmp_path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    t = Table(name="Padded")

    def cell(text=None, repeat=1):
        c = TableCell(numbercolumnsrepeated=repeat) if text is None else TableCell(valuetype="string",
                                                                                 numbercolumnsrepeated=repeat)
        if text is not None:
            c.addElement(P(text=text))
        return c
    r1 = TableRow()
    for c in (cell("a"), cell(None, 3), cell("b", 2), cell(None, 16000)):
        r1.addElement(c)
    t.addElement(r1)
    t.addElement(TableRow(numberrowsrepeated=2))
    r2 = TableRow(numberrowsrepeated=2)
    r2.addElement(cell("c"))
    t.addElement(r2)
    blank = TableRow(numberrowsrepeated=1_000_000)
    blank.addElement(cell(None, 1024))
    t.addElement(blank)
    doc.spreadsheet.addElement(t)
    src = tmp_path / "padded.ods"
    doc.save(str(src))
    rows = sheets.read_sheets(src)[0].rows
    assert rows == [["a", None, None, None, "b", "b"], [], [], ["c"], ["c"]]


def test_sheet_to_pdf_and_docx(tmp_path):
    src = make_workbook(tmp_path / "Book.xlsx")
    pdf = convert.convert(src, "sheet", "PDF")[0]
    text = pdf_text(pdf)
    assert pdf_pages(pdf) >= 1
    for s in ("Sales", "Ünïcode ✓", "Zürich", "日本語", "2024-03-01", "TRUE", "ключ"):
        assert s in text, s
    from docx import Document
    d = Document(convert.convert(src, "sheet", "DOCX")[0])
    assert len(d.tables) == 2
    assert [c.text for c in d.tables[0].rows[1].cells] == ["Zürich café", "3", "1.5", "2024-03-01", "TRUE"]
    assert d.tables[0].rows[3].cells[3].text == ""  # short rows are padded


def test_sheet_titles():
    assert sheets.sheet_titles(["a/b", "x" * 40, "Sales", "sales", "", "'q'"]) == \
        ["a_b", "x" * 31, "Sales", "sales (2)", "Sheet5", "q"]


def test_never_overwrites_existing_output(tmp_path):
    src = make_workbook(tmp_path / "Book.xlsx")
    taken = tmp_path / "Book.ods"
    taken.write_text("mine")
    out = convert.convert(src, "sheet", "ODS")[0]
    assert out.name == "Book (1).ods" and taken.read_text() == "mine"


# ---------------------------------------------------------------- PDF to spreadsheet
def test_pdf_tables_to_xlsx(tmp_path):
    import openpyxl
    out = convert.convert(make_table_pdf(tmp_path / "invoice.pdf"), "pdf", "XLSX")[0]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Page 1"]
    rows = [[c.value for c in r] for r in wb.active.iter_rows()]
    assert rows == [["Item", "Count", "Cost"], ["Apples", 12, 3.5], ["Pears", 7, 2.25]]


def test_pdf_without_tables_to_xlsx(tmp_path):
    import openpyxl
    out = convert.convert(make_text_pdf(tmp_path / "plain.pdf"), "pdf", "XLSX")[0]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Page 1", "Page 2"]
    assert wb["Page 2"]["A1"].value == "Hello page 2"
    scan = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page().insert_image(fitz.Rect(0, 0, 200, 120), stream=png_bytes())
    doc.save(scan)
    with pytest.raises(UserError, match="scan"):
        convert.convert(scan, "pdf", "XLSX")


# ---------------------------------------------------------------- slides
def test_pptx_to_pdf_renders_text_pictures_and_skips_hidden(tmp_path):
    src = make_deck(tmp_path / "Deck.pptx")
    before = digest(src)
    out = convert.convert(src, "slides", "PDF")[0]
    with fitz.open(out) as d:
        assert d.page_count == 3  # the hidden slide isn't exported
        assert (round(d[0].rect.width), round(d[0].rect.height)) == (720, 540)
        texts = [p.get_text() for p in d]
        assert "Quarterly Review" in texts[0] and "Grüße aus Zürich 日本" in texts[0]
        assert "First point" in texts[1] and "Sub point" in texts[1] and "•" in texts[1]
        assert "Oval" in texts[2] and "Header cell" in texts[2] and "Body cell" in texts[2]
        assert d[2].get_images(), "the picture should be embedded"
        assert "Secret slide" not in "".join(texts)
    assert digest(src) == before and not leftovers(tmp_path)


def test_pptx_to_images(tmp_path):
    src = make_deck(tmp_path / "Deck.pptx")
    pngs = convert.convert(src, "slides", "PNG")
    assert [p.name for p in pngs] == ["Deck slide 1.png", "Deck slide 2.png", "Deck slide 3.png"]
    im = Image.open(pngs[2]).convert("RGB")
    assert im.size == (2000, 1500)  # 10 x 7.5 in at the default 200 dpi
    r, g, b = im.getpixel((300, 300))  # inside the green picture at 1 in, 1 in
    assert abs(r - 30) < 10 and abs(g - 160) < 10 and abs(b - 90) < 10
    jpgs = convert.convert(src, "slides", "JPG")
    assert len(jpgs) == 3 and Image.open(jpgs[0]).format == "JPEG"
    single = make_deck(tmp_path / "One.pptx", hidden=False)
    from pptx import Presentation
    prs = Presentation(single)
    for sld in list(prs.slides._sldIdLst)[1:]:
        prs.slides._sldIdLst.remove(sld)
    prs.save(single)
    assert [p.name for p in convert.convert(single, "slides", "PNG")] == ["One.png"]


def test_slide_background_theme_and_groups(tmp_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    group = s.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Grouped text"
    # move the group 2 inches right without moving its children's own coordinates
    off = group._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}off")
    off.set("x", str(int(off.get("x")) + Inches(2)))
    src = tmp_path / "styled.pptx"
    prs.save(src)
    with slides.render(src) as d:
        page = d[0]
        pix = page.get_pixmap(dpi=72)
        assert pix.pixel(5, 5) == (0x11, 0x22, 0x33)
        hits = page.search_for("Grouped text")
        assert hits and hits[0].x0 > Emu(Inches(3)).pt - 2  # 1 in + 2 in shift, in points


def test_slide_errors(tmp_path):
    ppt = tmp_path / "old.ppt"
    ppt.write_bytes(sheets.OLE_MAGIC + bytes(1000))
    with pytest.raises(UserError, match=r"\.pptx"):
        convert.convert(ppt, "slides", "PDF")
    locked = tmp_path / "locked.pptx"
    locked.write_bytes(sheets.OLE_MAGIC + "EncryptionInfo".encode("utf-16-le"))
    with pytest.raises(UserError, match="password"):
        convert.convert(locked, "slides", "PNG")
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"PK\x03\x04not really")
    with pytest.raises(UserError, match="damaged"):
        convert.convert(broken, "slides", "PDF")
    assert not list(tmp_path.glob("*.pdf")) and not list(tmp_path.glob("*.png"))


@pytest.mark.skipif(sys.platform != "darwin", reason="Keynote is macOS-only")
def test_keynote_export_when_installed(tmp_path, monkeypatch):
    from dragshift import platform_mac
    monkeypatch.setattr(documents, "iwork_to_pdf", REAL_IWORK)
    monkeypatch.setattr(platform_mac, "iwork_app", lambda name: Path(f"/Applications/{name}.app"))
    src = make_deck(tmp_path / "-My deck ü.pptx")
    calls = []

    def fake_run(cmd, **kw):
        calls.append(cmd)
        doc = fitz.open()
        doc.new_page().insert_text((72, 72), "From Keynote")
        doc.save(cmd[-1])
        return subprocess.CompletedProcess(cmd, 0, b"", b"")
    monkeypatch.setattr(subprocess, "run", fake_run)
    out = convert.convert(src, "slides", "PDF")[0]
    assert "From Keynote" in pdf_text(out)
    cmd = calls[0]
    split = cmd.index("--")
    script = cmd[1:split]
    assert cmd[0] == "osascript" and script[0] == "-e" and any('"Keynote"' in s for s in script)
    assert not any(str(tmp_path) in s for s in script), "paths must be arguments, not script text"
    assert cmd[split + 1] == str(src.resolve()) and cmd[split + 2].endswith(".pdf")
    assert len(convert.convert(src, "slides", "PNG")) == 1
    assert not leftovers(tmp_path)

    def failing_run(cmd, **kw):
        return subprocess.CompletedProcess(cmd, 1, b"", b"not allowed")
    monkeypatch.setattr(subprocess, "run", failing_run)
    fallback = convert.convert(src, "slides", "PDF")[0]
    assert "Quarterly Review" in pdf_text(fallback) and not leftovers(tmp_path)


def test_pdf_to_pptx(tmp_path, monkeypatch):
    from pptx import Presentation
    src = make_text_pdf(tmp_path / "Doc.pdf", pages=3)
    out = convert.convert(src, "pdf", "PPTX")[0]
    prs = Presentation(out)
    assert len(prs.slides) == 3
    with fitz.open(src) as d:  # PyMuPDF's default page is A4
        size = (round(d[0].rect.width * 12700), round(d[0].rect.height * 12700))
    assert (prs.slide_width, prs.slide_height) == size
    for i, s in enumerate(prs.slides, 1):
        assert len(s.shapes) == 1 and s.shapes[0].shape_type == 13  # PICTURE
        assert f"Hello page {i}" in s.notes_slide.notes_text_frame.text
    monkeypatch.setattr(slides, "MAX_PDF_PAGES", 2)
    with pytest.raises(UserError, match="Split"):
        convert.convert(src, "pdf", "PPTX")
    locked = tmp_path / "locked.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(locked, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(UserError, match="password"):
        convert.convert(locked, "pdf", "PPTX")


# ---------------------------------------------------------------- Markdown, HTML, text
def test_markdown_to_html(markdown_file):
    page = convert.convert(markdown_file, "markdown", "HTML")[0].read_text(encoding="utf-8")
    assert page.startswith("<!DOCTYPE html>") and "<title>Report title</title>" in page
    assert "<h1>Report title</h1>" in page and "<strong>bold</strong>" in page
    assert '<a href="https://example.com/page">link</a>' in page
    assert '<code class="language-python">' in page and "<th>Name</th>" in page
    assert "<ul>" in page and "nested item" in page and "<blockquote>" in page
    assert "<script" not in page and 'alert("x")' not in page
    assert 'href="javascript' not in page


def test_markdown_to_docx(markdown_file):
    from docx import Document
    d = Document(convert.convert(markdown_file, "markdown", "DOCX")[0])
    styles = [(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]
    assert ("Heading 1", "Report title") in styles and ("Heading 1", "Second chapter") in styles
    assert ("List Bullet", "first") in styles and ("List Bullet 2", "nested item") in styles
    assert ("List Number", "two") in styles
    code = next(p for p in d.paragraphs if 'print("hi")' in p.text)
    assert code.runs[0].font.name == "Courier New"
    bold = next(r for p in d.paragraphs for r in p.runs if r.text == "bold")
    assert bold.bold
    assert d.tables[0].cell(1, 0).text == "Zürich"
    links = [r for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert [r.target_ref for r in links] == ["https://example.com/page"]  # the javascript: link was dropped
    assert d.core_properties.title == "Report title"


def test_markdown_to_pdf_txt_and_back(markdown_file):
    pdf = convert.convert(markdown_file, "markdown", "PDF")[0]
    text = pdf_text(pdf)
    assert "Report title" in text and "nested item" in text and "ünïcödé" in text and "**" not in text
    txt = convert.convert(markdown_file, "markdown", "TXT")[0].read_text(encoding="utf-8")
    assert "Report title" in txt and "link (https://example.com/page)" in txt and "**" not in txt
    assert "- first" in txt and "    - nested item" not in txt  # plain list, nesting by indent
    html_file = convert.convert(markdown_file, "markdown", "HTML")[0]
    md = convert.convert(html_file, "text", "MD")[0].read_text(encoding="utf-8")
    assert md.startswith("# Report title")
    assert "**bold**" in md and "*italic*" in md and "[link](https://example.com/page)" in md
    assert "- second\n  - nested item" in md and "1. one\n2. two" in md
    assert '```python\nprint("hi")\n```' in md and "| Name | Qty |" in md and "> quoted words" in md


def test_text_and_html_exports(tmp_path):
    t = tmp_path / "plain notes.txt"
    t.write_text("First para line one\nline two <b>not bold</b>\n\nSecond para\n", encoding="utf-8")
    page = convert.convert(t, "text", "HTML")[0].read_text(encoding="utf-8")
    assert "<p>First para line one<br>line two &lt;b&gt;not bold&lt;/b&gt;</p>" in page
    epub = convert.convert(t, "text", "EPUB")[0]
    assert "Second para" in pdf_text(epub)
    assert convert.convert(t, "text", "PDF")[0].exists()  # the original monospace path still works
    h = tmp_path / "page.html"
    h.write_text("<html><head><title>My Page</title><style>p{color:red}</style></head><body>"
                 "<h1>Big</h1><p onclick='x()'>Hello <em>there</em></p><script>evil()</script>"
                 f'<img src="data:image/png;base64,{__import__("base64").b64encode(png_bytes()).decode()}" alt="pic">'
                 "</body></html>", encoding="utf-8")
    assert convert.convert(h, "text", "TXT")[0].read_text(encoding="utf-8").startswith("Big\n\nHello there")
    assert "evil" not in convert.convert(h, "text", "MD")[0].read_text(encoding="utf-8")
    pdf = convert.convert(h, "text", "PDF")[0]
    assert "Hello there" in pdf_text(pdf) and "color:red" not in pdf_text(pdf)
    with zipfile.ZipFile(convert.convert(h, "text", "EPUB")[0]) as z:
        assert any(n.startswith("OEBPS/images/") for n in z.namelist())
        assert b"My Page" in z.read("OEBPS/content.opf")
    json_file = tmp_path / "data.json"
    json_file.write_text('{"a": 1}')
    assert "PDF" in convert.targets_for(category(json_file), ".json")
    assert convert.convert(json_file, "text", "DOCX")[0].exists()


def test_markdown_local_pictures_kept_remote_ones_never_fetched(tmp_path):
    (tmp_path / "blue.png").write_bytes(png_bytes((300, 200), (10, 20, 200)))
    src = tmp_path / "pics.md"
    src.write_text("# Pics\n\nBefore ![alt](blue.png) after ![remote](https://example.com/x.png) end.\n",
                   encoding="utf-8")
    with fitz.open(convert.convert(src, "markdown", "PDF")[0]) as d:
        assert len(d[0].get_images()) == 1 and "end." in d[0].get_text()
    from docx import Document
    assert len(Document(convert.convert(src, "markdown", "DOCX")[0]).inline_shapes) == 1
    page = convert.convert(src, "markdown", "HTML")[0].read_text(encoding="utf-8")
    assert page.count("data:image/png;base64,") == 1 and 'src="https://example.com/x.png"' in page


def test_pictures_outside_the_documents_folder_are_never_read(tmp_path):
    """A downloaded Markdown or HTML file can't pull in other local files (../, absolute paths,
    or a link that leads out of its folder); pictures in its own folder still work."""
    secret = tmp_path / "secret"
    secret.mkdir()
    (secret / "private.png").write_bytes(png_bytes((20, 20), (200, 0, 0)))
    folder = tmp_path / "dl"
    (folder / "img").mkdir(parents=True)
    (folder / "img" / "ok.png").write_bytes(png_bytes((20, 20), (0, 200, 0)))
    refs = ["../secret/private.png", str(secret / "private.png"), (secret / "private.png").as_uri(),
            "img/../../secret/private.png"]
    if sys.platform != "win32":
        (folder / "link.png").symlink_to(secret / "private.png")
        (folder / "linkdir").symlink_to(secret, target_is_directory=True)
        refs += ["link.png", "linkdir/private.png"]
    md = folder / "readme.md"
    md.write_text("# Hi\n\n" + "".join(f"![x]({r})\n\n" for r in refs) + "![ok](img/ok.png)\n", encoding="utf-8")
    page = convert.convert(md, "markdown", "HTML")[0].read_text(encoding="utf-8")
    assert page.count("data:image/png;base64,") == 1
    html_src = folder / "page.html"
    html_src.write_text("<h1>Hi</h1>" + "".join(f'<img src="{r}">' for r in refs), encoding="utf-8")
    with fitz.open(convert.convert(html_src, "text", "PDF")[0]) as d:
        assert not d[0].get_images()
    doc = documents.Doc("t", "", base=folder)
    for r in refs:
        assert documents._image_bytes(doc, r) is None, r
    assert documents._image_bytes(doc, "img/ok.png") is not None


# ---------------------------------------------------------------- EPUB
def test_markdown_to_epub_structure(tmp_path, markdown_file):
    (tmp_path / "pic.png").write_bytes(png_bytes())
    markdown_file.write_text(MARKDOWN + "\n![A picture](pic.png)\n", encoding="utf-8")
    out = convert.convert(markdown_file, "markdown", "EPUB")[0]
    with zipfile.ZipFile(out) as z:
        infos = z.infolist()
        assert infos[0].filename == "mimetype" and infos[0].compress_type == zipfile.ZIP_STORED
        assert z.read("mimetype") == b"application/epub+zip"
        opf = etree.fromstring(z.read("OEBPS/content.opf"))
        ns = {"o": "http://www.idpf.org/2007/opf", "dc": "http://purl.org/dc/elements/1.1/"}
        assert opf.findtext(".//dc:title", namespaces=ns) == "Report title"
        assert len(opf.findall(".//o:spine/o:itemref", ns)) == 2  # one chapter per h1
        assert opf.find(".//o:item[@properties='nav']", ns) is not None
        for name in z.namelist():
            if name.endswith((".xhtml", ".opf", ".ncx", ".xml")):
                etree.fromstring(z.read(name))  # every part is well-formed XML
        chapter2 = z.read("OEBPS/chapter2.xhtml").decode("utf-8")
        assert "Second chapter" in chapter2 and "ünïcödé" in chapter2
        assert any(n.startswith("OEBPS/images/image1.png") for n in z.namelist())
    with fitz.open(out) as book:
        assert "Report title" in book[0].get_text() or "Report title" in "".join(p.get_text() for p in book)


def test_epub_to_text_formats(tmp_path):
    src = make_epub(tmp_path / "Book.epub",
                    ["<h1>Chapter One</h1><p>It was a <b>dark</b> night in Zürich.</p>"
                     '<p><img src="../images/cover.png" alt="cover"/></p>',
                     "<h1>Chapter Two</h1><ul><li>alpha</li><li>beta</li></ul>"],
                    images={"images/cover.png": png_bytes()})
    before = digest(src)
    txt = convert.convert(src, "epub", "TXT")[0].read_text(encoding="utf-8")
    assert "Chapter One" in txt and "dark night in Zürich" in txt and "- beta" in txt
    md = convert.convert(src, "epub", "MD")[0].read_text(encoding="utf-8")
    assert "# Chapter One" in md and "**dark**" in md and "# Chapter Two" in md
    from docx import Document
    d = Document(convert.convert(src, "epub", "DOCX")[0])
    assert d.core_properties.title == "Test Book"
    assert any(p.text == "Chapter Two" and p.style.name == "Heading 1" for p in d.paragraphs)
    assert len(d.inline_shapes) == 1, "the EPUB's picture should be carried over"
    pdf = convert.convert(src, "epub", "PDF")[0]
    assert pdf_pages(pdf) >= 1 and "Chapter Two" in pdf_text(pdf)
    assert digest(src) == before and not leftovers(tmp_path)


def test_epub_errors(tmp_path):
    font_only = ('<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
                 'xmlns:enc="http://www.w3.org/2001/04/xmlenc#"><enc:EncryptedData>'
                 '<enc:EncryptionMethod Algorithm="http://www.idpf.org/2008/embedding"/>'
                 "</enc:EncryptedData></encryption>")
    ok = make_epub(tmp_path / "fonts.epub", ["<p>Readable</p>"], encryption=font_only)
    assert "Readable" in convert.convert(ok, "epub", "TXT")[0].read_text(encoding="utf-8")
    drm = make_epub(tmp_path / "drm.epub", ["<p>x</p>"],
                    encryption=font_only.replace("http://www.idpf.org/2008/embedding",
                                                 "http://www.w3.org/2001/04/xmlenc#aes128-cbc"))
    for target in ("TXT", "PDF"):
        with pytest.raises(UserError, match="DRM"):
            convert.convert(drm, "epub", target)
    junk = tmp_path / "junk.epub"
    junk.write_bytes(b"not a zip")
    with pytest.raises(UserError, match="damaged"):
        convert.convert(junk, "epub", "MD")
    no_container = tmp_path / "bare.epub"
    with zipfile.ZipFile(no_container, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
    with pytest.raises(UserError, match="damaged"):
        convert.convert(no_container, "epub", "DOCX")
    empty = make_epub(tmp_path / "empty.epub", ["<p>x</p>"], spine=False)
    with pytest.raises(UserError, match="no readable chapters"):
        convert.convert(empty, "epub", "TXT")


def test_epub_repeated_spine_entries_are_read_once(tmp_path):
    src = make_epub(tmp_path / "Loop.epub", ["<p>Only once</p>"], repeat=50)
    doc = documents.read_epub(src)
    assert doc.body.count("Only once") == 1


def test_epub_total_size_is_capped(tmp_path, monkeypatch):
    """Many chapters and pictures that are each under the per-file limit still add up."""
    monkeypatch.setattr(documents, "MAX_BOOK_BYTES", 12_000)
    pics = {f"images/p{i}.png": png_bytes((80, 80), (i, 0, 0)) + b"\0" * 4000 for i in range(4)}
    chapters = [f'<p>Chapter {i}</p><p><img src="../images/p{i}.png"/></p>' for i in range(4)]
    src = make_epub(tmp_path / "Big.epub", chapters, images=pics)
    with pytest.raises(UserError, match="too large"):
        documents.read_epub(src)
    monkeypatch.setattr(documents, "MAX_BOOK_BYTES", 10 ** 6)
    assert documents.read_epub(src).body.count("Chapter") == 4


# ---------------------------------------------------------------- Word and PDF to text formats
def test_docx_to_md_html_epub(docx_file):
    md = convert.convert(docx_file, "docx", "MD")[0].read_text(encoding="utf-8")
    assert md.startswith("# Word heading")
    assert "Plain and **bold words** [a link](https://example.org/)" in md
    assert "- Bullet one\n  - Bullet nested" in md and "1. Number one" in md
    assert "| T00 |" in md
    page = convert.convert(docx_file, "docx", "HTML")[0].read_text(encoding="utf-8")
    assert "<h1>Word heading</h1>" in page and 'src="data:image/png;base64,' in page
    epub = convert.convert(docx_file, "docx", "EPUB")[0]
    with zipfile.ZipFile(epub) as z:
        assert "OEBPS/images/image1.png" in z.namelist()


def numbered_docx(path: Path) -> Path:
    """Word's own way of numbering: 'List Paragraph' plus numPr pointing into the numbering part
    (level 0 decimal, level 1 lowerLetter, level 2 bullet), and a second list whose level 0 is
    overridden to upperRoman."""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    d = Document()
    numbering = d.part.numbering_part.element
    levels = "".join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/></w:lvl>'
                     for i, fmt in enumerate(["decimal", "lowerLetter", "bullet"]))
    numbering.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">{levels}</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="90"/>'
                               '<w:lvlOverride w:ilvl="0"><w:lvl w:ilvl="0"><w:numFmt w:val="upperRoman"/>'
                               '</w:lvl></w:lvlOverride></w:num>'))

    def item(text, num_id, level):
        p = d.add_paragraph(text, style="List Paragraph")
        p._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/>'
                                               f'<w:numId w:val="{num_id}"/></w:numPr>'))
    item("First", 90, 0)
    item("Sub a", 90, 1)
    item("Dot", 90, 2)
    item("Second", 90, 0)
    d.add_paragraph("Between")
    item("Roman", 91, 0)
    d.save(path)
    return path


def test_word_numbered_lists_stay_numbered(tmp_path):
    src = numbered_docx(tmp_path / "Numbered.docx")
    body = convert.docx_to_html(src)
    assert "<ol>\n<li>First</li>\n<ol>\n<li>Sub a</li>\n<ul>\n<li>Dot</li>\n</ul>\n</ol>\n<li>Second</li>\n</ol>" in body
    assert "<ol>\n<li>Roman</li>\n</ol>" in body
    md = convert.convert(src, "docx", "MD")[0].read_text(encoding="utf-8")
    assert "1. First" in md and "2. Second" in md and "- Dot" in md and "1. Roman" in md


@pytest.mark.parametrize("target", convert.DOCX_TARGETS)
def test_damaged_and_password_protected_docx_give_the_same_message_for_every_target(tmp_path, target):
    damaged = tmp_path / "damaged.docx"
    damaged.write_bytes(b"not a zip at all")
    with pytest.raises(UserError, match="damaged or isn't really a .docx"):
        convert.convert(damaged, "docx", target)
    locked = tmp_path / "locked.docx"
    locked.write_bytes(sheets.OLE_MAGIC + bytes(500) + "EncryptionInfo".encode("utf-16-le") + bytes(500))
    with pytest.raises(UserError, match="password-protected"):
        convert.convert(locked, "docx", target)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["damaged.docx", "locked.docx"]


def test_encrypted_office_detector(tmp_path):
    from dragshift import util
    locked = tmp_path / "a.xlsx"
    locked.write_bytes(util.OLE_MAGIC + "EncryptionInfo".encode("utf-16-le"))
    old = tmp_path / "old.xls"
    old.write_bytes(util.OLE_MAGIC + bytes(100))
    zipped = tmp_path / "z.docx"
    zipped.write_bytes(b"PK\x03\x04" + "EncryptionInfo".encode("utf-16-le"))
    assert util.encrypted_ole(locked) and not util.encrypted_ole(old) and not util.encrypted_ole(zipped)
    assert not util.encrypted_ole(tmp_path / "missing.docx")


def test_docx_to_pdf_keeps_lists_and_pictures(docx_file):
    pdf = convert.convert(docx_file, "docx", "PDF")[0]
    with fitz.open(pdf) as d:
        text = d[0].get_text()
        assert "Bullet nested" in text and "a link" in text
        assert d[0].get_images(), "pictures should now be kept"


def test_pdf_to_md_and_epub(tmp_path):
    src = make_text_pdf(tmp_path / "Doc.pdf")
    md = convert.convert(src, "pdf", "MD")[0].read_text(encoding="utf-8")
    assert "# Hello page 1" in md and "Body text line for extraction." in md
    epub = convert.convert(src, "pdf", "EPUB")[0]
    assert "Hello page 2" in pdf_text(epub)
    scan = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page().insert_image(fitz.Rect(0, 0, 200, 120), stream=png_bytes())
    doc.save(scan)
    with pytest.raises(UserError, match="no text"):
        convert.convert(scan, "pdf", "MD")
    with zipfile.ZipFile(convert.convert(scan, "pdf", "EPUB")[0]) as z:  # pictures alone still make a book
        assert any(n.startswith("OEBPS/images/") for n in z.namelist())


# ---------------------------------------------------------------- RTF
def test_rtf_without_textutil(tmp_path, monkeypatch):
    if sys.platform == "darwin":
        from dragshift import platform_mac
        monkeypatch.setattr(platform_mac, "textutil_html", lambda src: None)
    src = tmp_path / "letter.rtf"
    src.write_text(RTF, encoding="ascii")
    txt = convert.convert(src, "rtf", "TXT")[0].read_text(encoding="utf-8")
    assert "Grüße aus Zürich" in txt and "Second line here" in txt
    from docx import Document
    assert any("Grüße" in p.text for p in Document(convert.convert(src, "rtf", "DOCX")[0]).paragraphs)
    assert "Zürich" in pdf_text(convert.convert(src, "rtf", "PDF")[0])
    assert "Second line here" in convert.convert(src, "rtf", "MD")[0].read_text(encoding="utf-8")


@pytest.mark.skipif(sys.platform != "darwin", reason="textutil is macOS-only")
def test_rtf_with_textutil_keeps_bold(tmp_path):
    src = tmp_path / "-styled ü.rtf"
    src.write_text(RTF, encoding="ascii")
    md = convert.convert(src, "rtf", "MD")[0].read_text(encoding="utf-8")
    assert "Grüße aus **Zürich**" in md


def test_rtfd_bundle_and_rtf_errors(tmp_path, monkeypatch):
    bundle = tmp_path / "Note.rtfd"
    bundle.mkdir()
    (bundle / "TXT.rtf").write_text(RTF, encoding="ascii")
    assert "Grüße" in convert.convert(bundle, "rtf", "TXT")[0].read_text(encoding="utf-8")
    assert (tmp_path / "Note.txt").exists()
    empty_bundle = tmp_path / "Empty.rtfd"
    empty_bundle.mkdir()
    with pytest.raises(UserError, match="no text"):
        convert.convert(empty_bundle, "rtf", "TXT")
    fake = tmp_path / "fake.rtf"
    fake.write_text("just text")
    with pytest.raises(UserError, match="RTF"):
        convert.convert(fake, "rtf", "TXT")


# ---------------------------------------------------------------- Pages
def make_pages(path: Path, entries: dict) -> Path:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Index/Document.iwa", b"\x00binary")
        for name, data in entries.items():
            z.writestr(name, data)
    return path


def jpeg(size) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, (200, 40, 40)).save(buf, "JPEG")
    return buf.getvalue()


def test_pages_with_preview_pdf(tmp_path):
    preview = make_text_pdf(tmp_path / "preview-src.pdf", pages=2).read_bytes()
    src = make_pages(tmp_path / "Old Letter.pages", {"QuickLook/Preview.pdf": preview,
                                                       "QuickLook/Thumbnail.jpg": jpeg((60, 80))})
    out = convert.convert(src, "pages", "PDF")[0]
    assert out.name == "Old Letter.pdf" and pdf_pages(out) == 2 and "Hello page 2" in pdf_text(out)


def test_pages_with_preview_images_picks_the_largest(tmp_path):
    src = make_pages(tmp_path / "-new doc.pages", {"preview-micro.jpg": jpeg((40, 50)),
                                                    "preview.jpg": jpeg((800, 1000)),
                                                    "preview-web.jpg": jpeg((400, 500))})
    out = convert.convert(src, "pages", "PDF")[0]
    with fitz.open(out) as d:
        assert d.page_count == 1
        xref = d[0].get_images()[0][0]
        assert fitz.Pixmap(d, xref).width == 800


def test_pages_errors_and_bundle(tmp_path):
    bare = make_pages(tmp_path / "bare.pages", {})
    with pytest.raises(UserError, match="Export"):
        convert.convert(bare, "pages", "PDF")
    junk = tmp_path / "junk.pages"
    junk.write_bytes(b"nope")
    with pytest.raises(UserError, match="damaged"):
        convert.convert(junk, "pages", "PDF")
    bundle = tmp_path / "Folder.pages"
    (bundle / "QuickLook").mkdir(parents=True)
    (bundle / "QuickLook" / "Thumbnail.jpg").write_bytes(jpeg((120, 160)))
    assert pdf_pages(convert.convert(bundle, "pages", "PDF")[0]) == 1
    assert not list(tmp_path.glob("bare*.pdf")) and not list(tmp_path.glob("junk*.pdf"))


# ---------------------------------------------------------------- file names and summaries
@pytest.mark.parametrize("name", ["-leading dash.xlsx", "spaces and ünïcode 日本.xlsx", "--.xlsx"])
def test_awkward_file_names(tmp_path, name):
    src = make_workbook(tmp_path / name)
    before = digest(src)
    outs = convert.convert(src, "sheet", "CSV") + convert.convert(src, "sheet", "PDF")
    assert all(o.parent == tmp_path and o.exists() for o in outs)
    assert all(o.name.startswith(src.stem) for o in outs)
    assert digest(src) == before and not leftovers(tmp_path)


def test_result_summaries(tmp_path):
    book = make_workbook(tmp_path / "Book.xlsx")
    csvs = convert.convert(book, "sheet", "CSV")
    r = report.summarize("convert:CSV", [book], csvs)
    assert r.title == "2 sheets → 2 CSVs"
    ods = convert.convert(book, "sheet", "ODS")
    r = report.summarize("convert:ODS", [book], ods)
    assert r.title == "XLSX → ODS" and "2 sheets" in r.detail and "6 rows" in r.detail
    deck = make_deck(tmp_path / "Deck.pptx")
    pngs = convert.convert(deck, "slides", "PNG")
    assert report.summarize("convert:PNG", [deck], pngs).title == "3 slides → 3 PNGs"
    pptx = convert.convert(make_text_pdf(tmp_path / "Doc.pdf"), "pdf", "PPTX")
    assert "2 slides" in report.summarize("convert:PPTX", [tmp_path / "Doc.pdf"], pptx).detail
    pdf = convert.convert(deck, "slides", "PDF")
    assert "3 pages" in report.summarize("convert:PDF", [deck], pdf).detail
