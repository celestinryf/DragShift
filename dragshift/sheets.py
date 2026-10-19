"""Spreadsheets: XLSX, XLS (read only), ODS, CSV and TSV.

Every format is read into a list of `Sheet`s (a name plus rows of plain Python values) and
written back out from that list, so each new format only needs a reader and a writer."""
from __future__ import annotations

import csv
import datetime as dt
import html
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .util import CONTROL_CHARS, OLE_MAGIC, UserError, encrypted_ole

# Reading stops past this many cells: a few million Python values already use about a gigabyte.
MAX_CELLS = 2_000_000
# PDF and Word tables stop being readable (and take minutes to lay out) past this many rows.
MAX_TABLE_ROWS = 5_000

ZIP_MAGIC = b"PK\x03\x04"
ODS_MIME = b"application/vnd.oasis.opendocument.spreadsheet"
INT_RE = re.compile(r"-?(0|[1-9]\d{0,14})")          # no leading zeros: "007" stays text
FLOAT_RE = re.compile(r"-?(\d+\.\d*|\.\d+)([eE][-+]?\d+)?|-?\d+[eE][-+]?\d+")
BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
BAD_FILE_CHARS = re.compile(r'[\x00-\x1f<>:"/\\|?*]')

PASSWORD_MSG = "This spreadsheet is password-protected. Remove the password first, then try again."


@dataclass
class Sheet:
    name: str
    rows: list[list]
    from_text: bool = False   # values are raw CSV strings: numbers are recognised when writing XLSX/ODS


# ---------------------------------------------------------------- reading
def read_sheets(src: str | Path) -> list[Sheet]:
    """All sheets of a workbook (CSV/TSV give one). Raises UserError if nothing can be read."""
    src = Path(src)
    kind = _sniff(src)
    if kind == "xlsx":
        sheets = _read_xlsx(src)
    elif kind == "ods":
        sheets = _read_ods(src)
    elif kind == "xls":
        sheets = _read_xls(src)
    elif src.suffix.lower() in (".csv", ".tsv"):
        sheets = [_read_delimited(src)]
    else:
        raise UserError(f"This {src.suffix} file isn't a spreadsheet DragShift can read. "
                        "Open it in a spreadsheet app and save it as XLSX or CSV.")
    if not any(s.rows for s in sheets):
        raise UserError("This spreadsheet is empty.")
    return sheets


def _sniff(src: Path) -> str:
    """'xlsx', 'ods', 'xls' or 'text', from the file's contents (names can lie)."""
    with open(src, "rb") as f:
        head = f.read(8)
    if head.startswith(ZIP_MAGIC):
        try:
            with zipfile.ZipFile(src) as z:
                if "mimetype" in z.namelist() and z.read("mimetype").strip() == ODS_MIME:
                    return "ods"
        except zipfile.BadZipFile:
            raise UserError("This spreadsheet is damaged and can't be opened.") from None
        return "xlsx"
    if head == OLE_MAGIC:
        if encrypted_ole(src):
            raise UserError(PASSWORD_MSG)
        return "xls"
    return "text"


class _Budget:
    """Counts cells as they're read and stops when a sheet is too big to handle."""

    def __init__(self):
        self.cells = 0

    def take(self, n: int):
        self.cells += n
        if self.cells > MAX_CELLS:
            raise UserError(f"This spreadsheet is too big to convert (over {MAX_CELLS:,} cells).")


def trim(rows: list[list]) -> list[list]:
    """Drop empty cells at the end of each row and empty rows at the end of the sheet."""
    out = []
    for row in rows:
        row = list(row)
        while row and (row[-1] is None or row[-1] == ""):
            row.pop()
        out.append(row)
    while out and not out[-1]:
        out.pop()
    return out


def _read_xlsx(src: Path) -> list[Sheet]:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    f = open(src, "rb")  # a file object, so a workbook renamed to .xls or .csv still opens
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        f.close()
        raise UserError("This spreadsheet is damaged or isn't really an Excel file.") from e
    budget = _Budget()
    sheets = []
    try:
        for ws in wb.worksheets:
            # the stored size can be wrong (some apps write A1:A1, others a million rows); recount
            ws.reset_dimensions()
            rows = []
            for row in ws.iter_rows(values_only=True):
                budget.take(len(row))
                rows.append(list(row))
            sheets.append(Sheet(ws.title, trim(rows)))
    finally:
        wb.close()
        f.close()
    return sheets


def _read_xls(src: Path) -> list[Sheet]:
    import xlrd
    try:
        book = xlrd.open_workbook(src, on_demand=True)
    except xlrd.XLRDError as e:
        if "encrypt" in str(e).lower():
            raise UserError(PASSWORD_MSG) from e
        raise UserError("This .xls file is damaged or isn't really an Excel file.") from e
    except Exception as e:  # xlrd raises assorted errors on garbage input
        raise UserError("This .xls file is damaged or isn't really an Excel file.") from e
    budget = _Budget()
    sheets = []
    try:
        for i in range(book.nsheets):
            sh = book.sheet_by_index(i)
            budget.take(sh.nrows * sh.ncols)
            rows = [[_xls_value(sh.cell(r, c), book.datemode) for c in range(sh.ncols)] for r in range(sh.nrows)]
            sheets.append(Sheet(sh.name, trim(rows)))
            book.unload_sheet(i)
    finally:
        book.release_resources()
    return sheets


def _xls_value(cell, datemode: int):
    import xlrd
    t, v = cell.ctype, cell.value
    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if t == xlrd.XL_CELL_NUMBER:
        return int(v) if float(v).is_integer() else v
    if t == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(v, datemode)
        except Exception:
            return v
    if t == xlrd.XL_CELL_BOOLEAN:
        return bool(v)
    if t == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(v, "#ERROR")
    return v


def _read_ods(src: Path) -> list[Sheet]:
    from odf.opendocument import load
    from odf.table import Table, TableRow
    try:
        with zipfile.ZipFile(src) as z:
            if b"encryption-data" in z.read("META-INF/manifest.xml"):
                raise UserError(PASSWORD_MSG)
        doc = load(str(src))
    except UserError:
        raise
    except Exception as e:  # odfpy raises zip, XML and key errors on damaged files
        raise UserError("This spreadsheet is damaged or isn't really an ODS file.") from e
    budget = _Budget()
    sheets = []
    for table in doc.spreadsheet.getElementsByType(Table):
        rows: list[list] = []
        blank_rows = 0  # empty rows are only added once something follows them
        for row in table.getElementsByType(TableRow):
            cells = _ods_row(row, budget)
            repeat = int(row.getAttribute("numberrowsrepeated") or 1)
            if not cells:
                blank_rows += repeat
                continue
            budget.take(blank_rows + len(cells) * repeat)
            rows.extend([] for _ in range(blank_rows))
            rows.extend(list(cells) for _ in range(repeat))
            blank_rows = 0
        sheets.append(Sheet(table.getAttribute("name") or f"Sheet{len(sheets) + 1}", trim(rows)))
    return sheets


def _ods_row(row, budget: _Budget) -> list:
    """Cell values of one ODS row. Files pad rows with thousands of repeated empty cells, so
    empty cells are only added once a value follows them."""
    from odf.namespaces import TABLENS
    cells: list = []
    blanks = 0
    for cell in row.childNodes:
        if cell.qname not in ((TABLENS, "table-cell"), (TABLENS, "covered-table-cell")):
            continue
        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        value = _ods_value(cell)
        if value is None:
            blanks += repeat
            continue
        budget.take(blanks + repeat)
        cells.extend([None] * blanks)
        cells.extend([value] * repeat)
        blanks = 0
    return cells


def _ods_value(cell):
    from odf import teletype
    from odf.text import P
    vtype = cell.getAttribute("valuetype")
    try:
        if vtype in ("float", "percentage", "currency"):
            v = float(cell.getAttribute("value"))
            return int(v) if v.is_integer() and abs(v) < 1e15 else v
        if vtype == "boolean":
            return (cell.getAttribute("booleanvalue") or "").lower() == "true"
        if vtype == "date":
            return dt.datetime.fromisoformat(cell.getAttribute("datevalue"))
    except (TypeError, ValueError):
        pass  # fall back to the displayed text
    text = "\n".join(teletype.extractText(p) for p in cell.getElementsByType(P))
    return text if text else None


def _read_delimited(src: Path) -> Sheet:
    from .convert import read_text
    text = read_text(src)
    delim = "\t" if src.suffix.lower() == ".tsv" else _sniff_delimiter(text)
    budget = _Budget()
    rows = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delim):
            budget.take(len(row))
            rows.append(row)
    except csv.Error as e:
        raise UserError(f"This file isn't valid {src.suffix.lstrip('.').upper()}: {e}") from e
    return Sheet(src.stem, trim(rows), from_text=True)


def _sniff_delimiter(text: str) -> str:
    """Comma, or semicolon/tab/pipe when the file clearly uses one (common in Europe)."""
    sample = text[:65536]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


# ---------------------------------------------------------------- values
def typed(value):
    """A CSV string as a number when it plainly is one (like Excel does), else unchanged."""
    if not isinstance(value, str):
        return value
    s = value.strip()
    if not s:
        return None
    if INT_RE.fullmatch(s):
        return int(s)
    if FLOAT_RE.fullmatch(s):
        return float(s)
    return value


def cell_text(value) -> str:
    """How a value reads in CSV, TSV, PDF or Word."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _out_value(value, from_text: bool):
    if from_text:
        value = typed(value)
    if isinstance(value, str):
        return CONTROL_CHARS.sub("", value)
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)  # Excel has no time zones
    return value


def sheet_titles(names: list[str]) -> list[str]:
    """Valid, unique Excel sheet names: at most 31 characters, none of []:*?/\\ ."""
    used: set[str] = set()
    out = []
    for i, name in enumerate(names):
        base = BAD_SHEET_CHARS.sub("_", CONTROL_CHARS.sub("", str(name))).strip().strip("'")[:31] or f"Sheet{i + 1}"
        title, n = base, 2
        while title.lower() in used:
            suffix = f" ({n})"
            title = base[:31 - len(suffix)] + suffix
            n += 1
        used.add(title.lower())
        out.append(title)
    return out


def file_label(name: str) -> str:
    """A sheet name that's safe inside a file name on every platform."""
    return BAD_FILE_CHARS.sub("_", name).strip(" .") or "Sheet"


# ---------------------------------------------------------------- writing
def write_xlsx(sheets: list[Sheet], out: Path) -> None:
    import openpyxl
    from .convert import atomic_save
    wb = openpyxl.Workbook(write_only=True)
    for sheet, title in zip(sheets, sheet_titles([s.name for s in sheets])):
        ws = wb.create_sheet(title)
        for row in sheet.rows:
            ws.append([_xlsx_cell(ws, _out_value(v, sheet.from_text)) for v in row])
    atomic_save(wb.save, out)


def _xlsx_cell(ws, value):
    """The value as openpyxl should store it. openpyxl turns any string starting with "=" into a
    live formula, so a CSV cell like =HYPERLINK(...) would run when the file is opened in Excel:
    those are written as plain text cells instead."""
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    from openpyxl.cell import WriteOnlyCell
    cell = WriteOnlyCell(ws, value=value)
    cell.data_type = "s"
    return cell


def write_ods(sheets: list[Sheet], out: Path) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableColumn, TableRow
    from .convert import atomic_save
    doc = OpenDocumentSpreadsheet()
    for sheet, title in zip(sheets, sheet_titles([s.name for s in sheets])):
        table = Table(name=title)
        width = max((len(r) for r in sheet.rows), default=0)
        table.addElement(TableColumn(numbercolumnsrepeated=max(width, 1)))
        for row in sheet.rows:
            tr = TableRow()
            for v in row:
                tr.addElement(_ods_cell(_out_value(v, sheet.from_text)))
            table.addElement(tr)
        doc.spreadsheet.addElement(table)
    atomic_save(lambda dst: doc.save(str(dst)), out)


def _ods_cell(v):
    from odf import teletype
    from odf.table import TableCell
    from odf.text import P
    if v is None:
        return TableCell()
    if isinstance(v, bool):
        cell = TableCell(valuetype="boolean", booleanvalue="true" if v else "false")
    elif isinstance(v, (int, float)):
        cell = TableCell(valuetype="float", value=repr(v))
    elif isinstance(v, (dt.datetime, dt.date)):
        cell = TableCell(valuetype="date", datevalue=v.isoformat())
    else:
        cell = TableCell(valuetype="string")
    for line in cell_text(v).split("\n"):
        p = P()
        teletype.addTextToElement(p, line)
        cell.addElement(p)
    return cell


def write_delimited(src: Path, sheets: list[Sheet], ext: str) -> list[Path]:
    """One CSV/TSV per non-empty sheet: 'Book.csv', or 'Book - Sales.csv' for several sheets."""
    from .convert import atomic_save
    from .util import sibling
    delim = "\t" if ext == ".tsv" else ","
    real = [s for s in sheets if s.rows]
    outs = []
    for sheet in real:
        out = sibling(src, f"- {file_label(sheet.name)}" if len(real) > 1 else "", ext)

        def write(dst, sheet=sheet):
            # UTF-8 with a byte-order mark: without it Excel reads accented letters as gibberish
            with open(dst, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f, delimiter=delim)
                for row in sheet.rows:
                    w.writerow([cell_text(v) for v in row])
        atomic_save(write, out)
        outs.append(out)
    return outs


def sheets_html(sheets: list[Sheet]) -> str:
    """Non-empty sheets as HTML tables (first row as the header), for PDF and Word export."""
    real = [s for s in sheets if s.rows]
    total = sum(len(s.rows) for s in real)
    if total > MAX_TABLE_ROWS:
        raise UserError(f"This spreadsheet has {total:,} rows, too many for a readable document "
                        f"(the limit is {MAX_TABLE_ROWS:,}). Convert it to CSV or XLSX instead.")
    parts = []
    for sheet in real:
        if len(real) > 1:
            parts.append(f"<h2>{html.escape(sheet.name)}</h2>")
        width = max(len(r) for r in sheet.rows)
        rows = []
        for i, row in enumerate(sheet.rows):
            tag = "th" if i == 0 else "td"
            cells = [html.escape(cell_text(v)) for v in row] + [""] * (width - len(row))
            rows.append("<tr>" + "".join(f"<{tag}>{c}</{tag}>" for c in cells) + "</tr>")
        parts.append("<table>" + "".join(rows) + "</table>")
    return "\n".join(parts)


def row_count(sheets: list[Sheet]) -> int:
    return sum(len(s.rows) for s in sheets)
