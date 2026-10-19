"""Slides: PowerPoint to PDF, JPG and PNG, and PDF to PowerPoint.

On a Mac with Keynote installed, decks are exported through Keynote (best fidelity, and it also
reads old .ppt files). Everywhere else a built-in renderer draws each slide into a PDF page with
PyMuPDF: backgrounds, pictures, text with its inherited sizes, colours and bullets, tables,
lines and simple shapes. Charts, SmartArt, gradients, shadows, rotation and unusual shapes are
simplified. Slide images are made by rendering those PDF pages."""
from __future__ import annotations

import colorsys
import contextlib
import html
import io
import traceback
from pathlib import Path

from PIL import Image

from .convert import atomic_save, fitz, open_pdf, save_image
from .util import CONFIG, CONTROL_CHARS, OLE_MAGIC, UserError, encrypted_ole, output_dir, sibling

EMU_PER_PT = 12700
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
DEFAULT_SCHEME = {"dk1": (0, 0, 0), "lt1": (255, 255, 255), "dk2": (68, 84, 106), "lt2": (231, 230, 230),
                  "accent1": (68, 114, 196), "accent2": (237, 125, 49), "accent3": (165, 165, 165),
                  "accent4": (255, 192, 0), "accent5": (91, 155, 213), "accent6": (112, 173, 71),
                  "hlink": (5, 99, 193), "folHlink": (149, 79, 114)}
PRESET_COLORS = {"black": (0, 0, 0), "white": (255, 255, 255), "red": (255, 0, 0), "green": (0, 128, 0),
                 "blue": (0, 0, 255), "yellow": (255, 255, 0), "gray": (128, 128, 128), "grey": (128, 128, 128)}
DEFAULT_SIZE_PT = 18.0     # PowerPoint's size for text with no size set anywhere
DEFAULT_INSETS = (91440, 45720, 91440, 45720)  # left, top, right, bottom text margins (EMU)
MAX_PDF_PAGES = 500        # PDF -> PPTX: each page becomes a full-slide picture
PPTX_MIN_EMU, PPTX_MAX_EMU = 914400, 51206400   # PowerPoint's slide size limits (1 to 56 inches)
DEFAULT_TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"  # "Medium Style 2 - Accent 1"
TEXT_CSS = "* { font-family: sans-serif; } p { margin: 0; line-height: 1.12; }"


def _q(tag: str) -> str:
    prefix, name = tag.split(":")
    return f"{{{NS[prefix]}}}{name}"


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


# ---------------------------------------------------------------- entry points
def slides_to(src: str | Path, target: str) -> list[Path]:
    src = Path(src)
    target = target.upper()
    if target == "PDF":
        out = sibling(src, ext=".pdf")
        from .documents import iwork_to_pdf
        if iwork_to_pdf("Keynote", src, out):
            return [out]
        doc = render(src)
        atomic_save(lambda dst: doc.save(str(dst), garbage=3, deflate=True), out)
        doc.close()
        return [out]
    dpi = int(CONFIG.get("pdf_dpi", 200))
    outs = []
    with _deck_pdf(src) as doc:
        n = doc.page_count
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            out = sibling(src, f"slide {i + 1}" if n > 1 else "", "." + target.lower())
            save_image(im, out, target)
            outs.append(out)
    return outs


@contextlib.contextmanager
def _deck_pdf(src: Path):
    """The deck as an open PDF: exported by Keynote when possible, else rendered here."""
    from .documents import iwork_to_pdf
    tmp = output_dir(src) / f".~{src.stem}.slides.pdf"
    try:
        if iwork_to_pdf("Keynote", src, tmp):
            doc = fitz.open(tmp)
        else:
            doc = render(src)
        try:
            yield doc
        finally:
            doc.close()
    finally:
        with contextlib.suppress(OSError):
            tmp.unlink()


def open_pptx(src: Path):
    from pptx import Presentation
    with open(src, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        if encrypted_ole(src):
            raise UserError("This presentation is password-protected. Remove the password first, then try again.")
        raise UserError("Old PowerPoint .ppt files need Keynote (on a Mac) or PowerPoint to read. "
                        "Save it as .pptx first, then convert that.")
    try:
        return Presentation(str(src))
    except Exception as e:  # python-pptx raises zip, XML and key errors on damaged files
        raise UserError("This presentation is damaged or isn't really a PowerPoint file.") from e


def slide_count(src: str | Path) -> int:
    from pptx import Presentation
    return sum(1 for s in Presentation(str(src)).slides if s._element.get("show") != "0")


# ---------------------------------------------------------------- the renderer
def render(src: str | Path):
    """Draw every visible slide of a .pptx onto a page of a new in-memory PDF."""
    prs = open_pptx(Path(src))
    width = (prs.slide_width or 9144000) / EMU_PER_PT
    height = (prs.slide_height or 6858000) / EMU_PER_PT
    doc = fitz.open()
    scratch = fitz.open()  # text is laid out here first to measure it for vertical centring
    try:
        for slide in prs.slides:
            if slide._element.get("show") == "0":  # hidden slides aren't exported by PowerPoint either
                continue
            page = doc.new_page(width=width, height=height)
            ctx = _Context(prs, slide, scratch.new_page(width=width, height=height))
            ctx.draw_slide(page)
            scratch.delete_page(0)
    finally:
        scratch.close()
    if doc.page_count == 0:
        doc.close()
        raise UserError("This presentation has no visible slides.")
    return doc


class _Transform:
    """Maps shape coordinates (EMU, possibly inside groups) to page points."""

    def __init__(self, sx=1 / EMU_PER_PT, sy=1 / EMU_PER_PT, dx=0.0, dy=0.0):
        self.sx, self.sy, self.dx, self.dy = sx, sy, dx, dy

    def rect(self, x, y, w, h):
        x0, y0 = x * self.sx + self.dx, y * self.sy + self.dy
        return fitz.Rect(x0, y0, x0 + w * self.sx, y0 + h * self.sy)

    def point(self, x, y):
        return fitz.Point(x * self.sx + self.dx, y * self.sy + self.dy)

    def group(self, xfrm):
        """The transform for the children of a group with this a:xfrm."""
        def pair(tag, a, b):
            el = xfrm.find(_q(tag))
            return (int(el.get(a, 0)), int(el.get(b, 0))) if el is not None else (0, 0)
        off, ext = pair("a:off", "x", "y"), pair("a:ext", "cx", "cy")
        ch_off, ch_ext = pair("a:chOff", "x", "y"), pair("a:chExt", "cx", "cy")
        kx = ext[0] / ch_ext[0] if ch_ext[0] else 1.0
        ky = ext[1] / ch_ext[1] if ch_ext[1] else 1.0
        return _Transform(self.sx * kx, self.sy * ky,
                          self.sx * (off[0] - ch_off[0] * kx) + self.dx,
                          self.sy * (off[1] - ch_off[1] * ky) + self.dy)


class _Context:
    """Per-slide state: theme colours and fonts, the master's text styles, and a scratch page."""

    def __init__(self, prs, slide, scratch_page):
        self.prs, self.slide, self.scratch = prs, slide, scratch_page
        self.layout = slide.slide_layout
        self.master = self.layout.slide_master
        self.scheme = dict(DEFAULT_SCHEME)
        self.fonts = {"major": "", "minor": ""}
        self._load_theme()
        clr_map = self.master._element.find(_q("p:clrMap"))
        self.clr_map = dict(clr_map.attrib) if clr_map is not None else {}
        self.tx_styles = self.master._element.find(_q("p:txStyles"))
        self.default_style = prs.part._element.find(_q("p:defaultTextStyle"))

    def _load_theme(self):
        try:
            from lxml import etree
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            part = self.master.part.part_related_by(RT.THEME)
            root = etree.fromstring(part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
        except Exception:
            return
        scheme = root.find(".//a:clrScheme", NS)
        for el in scheme if scheme is not None else []:
            if len(el):
                c = el[0]
                val = c.get("val") if _local(c.tag) == "srgbClr" else c.get("lastClr")
                rgb = _hex(val)
                if rgb:
                    self.scheme[_local(el.tag)] = rgb
        for kind in ("major", "minor"):
            latin = root.find(f".//a:{kind}Font/a:latin", NS)
            if latin is not None:
                self.fonts[kind] = latin.get("typeface", "")

    # ------------------------------------------------------------ colours
    def color(self, el):
        """(r, g, b) floats 0-1 for a colour element (srgbClr, schemeClr...), or None."""
        if el is None:
            return None
        tag = _local(el.tag)
        if tag == "srgbClr":
            rgb = _hex(el.get("val"))
        elif tag == "schemeClr":
            name = el.get("val", "")
            name = self.clr_map.get(name, {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}.get(name, name))
            rgb = self.scheme.get(name)
        elif tag == "sysClr":
            rgb = _hex(el.get("lastClr")) or (0, 0, 0)
        elif tag == "prstClr":
            rgb = PRESET_COLORS.get(el.get("val", ""))
        else:
            return None
        if rgb is None:
            return None
        return _apply_mods(rgb, el)

    def fill_color(self, holder):
        """Colour of the first solidFill/gradFill under `holder` (spPr, rPr, tcPr...);
        False for an explicit noFill; None when nothing is set."""
        if holder is None:
            return None
        for child in holder:
            tag = _local(child.tag)
            if tag == "noFill":
                return False
            if tag == "solidFill" and len(child):
                return self.color(child[0])
            if tag == "gradFill":
                stop = child.find(".//a:gs", NS)
                if stop is not None and len(stop):
                    return self.color(stop[0])
        return None

    # ------------------------------------------------------------ drawing
    def draw_slide(self, page):
        self._draw_background(page)
        show_master = self.slide._element.get("showMasterSp") != "0"
        if show_master and self.layout._element.get("showMasterSp") != "0":
            self._draw_shapes(page, [s for s in self.master.shapes if not s.is_placeholder], _Transform())
        if show_master:
            self._draw_shapes(page, [s for s in self.layout.shapes if not s.is_placeholder], _Transform())
        self._draw_shapes(page, list(self.slide.shapes), _Transform())

    def _draw_background(self, page):
        for owner in (self.slide, self.layout, self.master):
            bg = owner._element.find("./p:cSld/p:bg", NS)
            if bg is None:
                continue
            pr = bg.find(_q("p:bgPr"))
            if pr is not None:
                blip = pr.find("./a:blipFill/a:blip", NS)
                if blip is not None and self._draw_blip(page, page.rect, owner.part, blip):
                    return
                color = self.fill_color(pr)
            else:
                ref = bg.find(_q("p:bgRef"))
                color = self.color(ref[0]) if ref is not None and len(ref) else None
            if color:
                page.draw_rect(page.rect, color=None, fill=color, width=0)
            return

    def _draw_shapes(self, page, shapes, tf):
        for shape in shapes:
            try:
                self._draw_shape(page, shape, tf)
            except Exception:  # one odd shape shouldn't cost the whole deck
                traceback.print_exc()

    def _draw_shape(self, page, shape, tf):
        el = shape._element
        tag = _local(el.tag)
        if tag == "grpSp":
            xfrm = el.find("./p:grpSpPr/a:xfrm", NS)
            self._draw_shapes(page, list(shape.shapes), tf.group(xfrm) if xfrm is not None else tf)
            return
        if shape.left is None or shape.width is None:
            return
        rect = tf.rect(shape.left, shape.top, shape.width, shape.height)
        if tag == "pic":
            blip = el.find("./p:blipFill/a:blip", NS)
            if blip is not None:
                crop = el.find("./p:blipFill/a:srcRect", NS)
                if not self._draw_blip(page, rect, shape.part, blip, crop):
                    page.draw_rect(rect, color=(0.75, 0.75, 0.75), fill=(0.95, 0.95, 0.95), width=0.5)
            return
        if tag == "graphicFrame":
            if getattr(shape, "has_table", False) and shape.has_table:
                self._draw_table(page, shape, tf)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                page.draw_rect(rect, color=(0.8, 0.8, 0.8), fill=(0.96, 0.96, 0.96), width=0.5)
                self._text(page, rect, '<p style="text-align:center;color:#888;font-size:12pt">Chart</p>', "ctr")
            return
        sppr = el.find(_q("p:spPr"))
        geom = el.find("./p:spPr/a:prstGeom", NS)
        preset = geom.get("prst", "rect") if geom is not None else "rect"
        if tag == "cxnSp" or preset in ("line", "straightConnector1"):
            self._draw_line(page, shape, tf, sppr)
            return
        if tag == "sp":
            self._draw_geometry(page, rect, preset, shape, sppr)
            self._draw_text(page, shape, rect)

    def _draw_blip(self, page, rect, part, blip, crop=None) -> bool:
        rid = blip.get(_q("r:embed"))
        if not rid:
            return False
        try:
            data = part.related_part(rid).blob
        except Exception:
            return False
        if crop is not None and any(crop.get(k) for k in ("l", "t", "r", "b")):
            data = _crop_image(data, crop)
        try:
            page.insert_image(rect, stream=data, keep_proportion=False)
            return True
        except Exception:
            pass
        try:  # formats MuPDF can't read but Pillow can (TIFF variants, WMF on Windows...)
            buf = io.BytesIO()
            Image.open(io.BytesIO(data)).convert("RGBA").save(buf, "PNG")
            page.insert_image(rect, stream=buf.getvalue(), keep_proportion=False)
            return True
        except Exception:
            return False

    def _style_color(self, shape, ref: str):
        """Colour from the shape's theme style (p:style/a:fillRef or a:lnRef)."""
        el = shape._element.find(f"./p:style/a:{ref}", NS)
        if el is None or el.get("idx") == "0" or not len(el):
            return None
        return self.color(el[0])

    def _outline(self, shape, sppr):
        """(colour, width in points) of a shape's outline, or (None, 0)."""
        ln = sppr.find(_q("a:ln")) if sppr is not None else None
        color = self.fill_color(ln) if ln is not None else None
        if color is False:
            return None, 0
        if color is None:
            color = self._style_color(shape, "lnRef")
        if not color:
            return None, 0
        w = int(ln.get("w", 0)) / EMU_PER_PT if ln is not None and ln.get("w") else 0.75
        return color, max(w, 0.25)

    def _draw_geometry(self, page, rect, preset, shape, sppr):
        fill = self.fill_color(sppr)
        if fill is None:
            fill = self._style_color(shape, "fillRef")
        stroke, width = self._outline(shape, sppr)
        if not fill and not stroke:
            return
        opacity = _alpha(sppr.find("./a:solidFill/*", NS)) if sppr is not None else 1.0
        s = page.new_shape()
        if preset == "ellipse":
            s.draw_oval(rect)
        elif preset == "roundRect":
            s.draw_rect(rect, radius=0.1)
        elif preset in ("triangle", "rtTriangle"):
            top = rect.tl if preset == "rtTriangle" else fitz.Point((rect.x0 + rect.x1) / 2, rect.y0)
            s.draw_polyline([top, rect.br, rect.bl, top])
        elif preset == "diamond":
            mx, my = (rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2
            pts = [fitz.Point(mx, rect.y0), fitz.Point(rect.x1, my), fitz.Point(mx, rect.y1), fitz.Point(rect.x0, my)]
            s.draw_polyline(pts + pts[:1])
        else:  # rectangles, and a box for every shape this renderer doesn't know
            s.draw_rect(rect)
        s.finish(color=stroke or None, fill=fill or None, width=width if stroke else 0,
                 fill_opacity=opacity, closePath=True)
        s.commit()

    def _draw_line(self, page, shape, tf, sppr):
        color, width = self._outline(shape, sppr)
        if not color:
            color, width = (0, 0, 0), 0.75
        xfrm = shape._element.find(".//a:xfrm", NS)
        flip_h = xfrm is not None and xfrm.get("flipH") == "1"
        flip_v = xfrm is not None and xfrm.get("flipV") == "1"
        x0, y0, x1, y1 = shape.left, shape.top, shape.left + shape.width, shape.top + shape.height
        if flip_h:
            x0, x1 = x1, x0
        if flip_v:
            y0, y1 = y1, y0
        page.draw_line(tf.point(x0, y0), tf.point(x1, y1), color=color, width=width)

    def _draw_table(self, page, shape, tf):
        table = shape.table
        widths = [c.width for c in table.columns]
        heights = [r.height for r in table.rows]
        tbl_pr = shape._element.find(".//a:tblPr", NS)
        styled = tbl_pr is not None and tbl_pr.findtext(_q("a:tableStyleId")) == DEFAULT_TABLE_STYLE
        header = styled and tbl_pr.get("firstRow") == "1"
        banded = styled and tbl_pr.get("bandRow") == "1"
        accent = tuple(c / 255 for c in self.scheme["accent1"])
        light = tuple(c / 255 for c in self.scheme.get(self.clr_map.get("bg1", "lt1"), (255, 255, 255)))
        y = shape.top
        for ri, row in enumerate(table.rows):
            x = shape.left
            is_header = header and ri == 0
            # PowerPoint's default style: accent header with light bold text, tinted bands below
            band = (ri - header) % 2 == 0
            style_fill = accent if is_header else (_tint(accent, 0.4 if band else 0.2) if banded else None)
            for ci, cell in enumerate(row.cells):
                if not cell.is_spanned:
                    w = sum(widths[ci:ci + max(cell.span_width, 1)])
                    h = sum(heights[ri:ri + max(cell.span_height, 1)])
                    rect = tf.rect(x, y, w, h)
                    tc_pr = cell._tc.find("./a:tcPr", NS)
                    fill = self.fill_color(tc_pr)
                    if fill is None:
                        fill = style_fill
                    border = (1, 1, 1) if styled else (0.6, 0.6, 0.6)
                    page.draw_rect(rect, color=border, fill=fill or None, width=0.75 if styled else 0.5)
                    body = cell._tc.find(_q("a:txBody"))
                    if body is not None:
                        inner = rect + (7.2, 3.6, -7.2, -3.6)
                        markup = self._paragraphs_html(body, [], None, 1.0, light if is_header else None, is_header)
                        self._text(page, inner, markup, _anchor_of(tc_pr) or "t")
                x += widths[ci]
            y += heights[ri]

    # ------------------------------------------------------------ text
    def _placeholder_chain(self, shape) -> list:
        """The shape's element, then the layout and master placeholders it inherits from."""
        chain = [shape._element]
        ph = shape
        for _ in range(2):
            try:
                ph = ph._base_placeholder if ph.is_placeholder else None
            except Exception:
                ph = None
            if ph is None:
                break
            chain.append(ph._element)
        return chain

    def _master_style(self, shape):
        if not shape.is_placeholder or self.tx_styles is None:
            return None
        kind = shape._element.find(".//p:nvPr/p:ph", NS).get("type", "obj")
        name = ("p:titleStyle" if kind in ("title", "ctrTitle")
                else "p:bodyStyle" if kind in ("body", "subTitle", "obj") else "p:otherStyle")
        return self.tx_styles.find(_q(name))

    def _draw_text(self, page, shape, rect):
        body = shape._element.find(_q("p:txBody"))
        if body is None or not "".join(body.itertext()).strip():
            return
        chain = self._placeholder_chain(shape)
        bodies = [e.find(_q("p:txBody")) for e in chain]
        body_prs = [b.find(_q("a:bodyPr")) for b in bodies if b is not None]
        insets = [int(_first_attr(body_prs, k, d)) for k, d in zip(("lIns", "tIns", "rIns", "bIns"), DEFAULT_INSETS)]
        left, top, right, bottom = (i / EMU_PER_PT for i in insets)
        inner = rect + (left, top, -right, -bottom)
        own = body.find(_q("a:bodyPr"))
        if own is not None and own.find(_q("a:spAutoFit")) is not None:
            inner.y1 = page.rect.y1  # the box grows with its text
        if inner.is_empty or inner.width < 2 or inner.height < 2:
            return
        fit = own.find(_q("a:normAutofit")) if own is not None else None
        scale = int(fit.get("fontScale", 100000)) / 100000 if fit is not None else 1.0
        anchor = None
        for bp in body_prs:
            anchor = _anchor_of(bp)
            if anchor:
                break
        markup = self._paragraphs_html(body, [b for b in bodies[1:] if b is not None], self._master_style(shape),
                                       scale, self._style_color(shape, "fontRef"))
        self._text(page, inner, markup, anchor or "t")

    def _text(self, page, rect, markup: str, anchor: str):
        if anchor in ("ctr", "b"):
            spare, _ = self.scratch.insert_htmlbox(rect, markup, css=TEXT_CSS, scale_low=0)
            if spare > 0:
                shift = spare / 2 if anchor == "ctr" else spare
                rect = fitz.Rect(rect.x0, rect.y0 + shift, rect.x1, rect.y1 + shift)
        page.insert_htmlbox(rect, markup, css=TEXT_CSS, scale_low=0)

    def _paragraphs_html(self, body, inherited_bodies: list, master_style, scale: float,
                         color=None, bold: bool = False) -> str:
        """HTML for a text body. `color` and `bold` are the defaults when nothing in the
        inheritance chain sets them (a shape's theme style, a table header)."""
        own_list = body.find(_q("a:lstStyle"))
        list_styles = [own_list] + [b.find(_q("a:lstStyle")) for b in inherited_bodies] + [master_style]
        list_styles = [s for s in list_styles if s is not None]
        out = []
        counters: dict[int, int] = {}
        for p in body.findall(_q("a:p")):
            ppr = p.find(_q("a:pPr"))
            level = int(ppr.get("lvl", 0)) if ppr is not None else 0
            name = _q(f"a:lvl{level + 1}pPr")
            levels = ([ppr] if ppr is not None else []) + [s.find(name) for s in list_styles]
            levels = [lv for lv in levels if lv is not None]
            defaults = [d for d in (lv.find(_q("a:defRPr")) for lv in levels) if d is not None]
            # the presentation-wide defaults come last: a shape's theme colour beats them
            base = self.default_style.find(name) if self.default_style is not None else None
            fallback = [base.find(_q("a:defRPr"))] if base is not None else []
            fallback = [d for d in fallback if d is not None]
            if base is not None:
                levels.append(base)
            align = {"ctr": "center", "r": "right", "just": "justify"}.get(_first_attr(levels, "algn", "l"), "left")
            spans = []
            size_used = None
            for child in p:
                tag = _local(child.tag)
                if tag in ("r", "fld"):
                    text = child.findtext(_q("a:t")) or ""
                    if text:
                        rpr = child.find(_q("a:rPr"))
                        style, size_used = self._run_css(rpr, defaults, fallback, scale, color, bold)
                        spans.append(f'<span style="{style}">{html.escape(text)}</span>')
                elif tag == "br":
                    spans.append("<br/>")
            end = p.find(_q("a:endParaRPr"))
            if size_used is None:
                _, size_used = self._run_css(end, defaults, fallback, scale, color, bold)
            bullet = self._bullet(levels, counters, level) if spans else ""
            indent = level * 18
            style = f"text-align:{align};font-size:{size_used:.1f}pt;padding-left:{indent}pt"
            content = "".join(spans) or "&nbsp;"
            out.append(f'<p style="{style}">{html.escape(bullet)}{content}</p>')
        return "".join(out)

    def _first_color(self, sources):
        for src in sources:
            color = self.fill_color(src)
            if color is not None:
                return color
        return None

    def _bullet(self, levels, counters, level) -> str:
        for lv in levels:
            if lv.find(_q("a:buNone")) is not None:
                return ""
            ch = lv.find(_q("a:buChar"))
            if ch is not None:
                return ch.get("char", "•") + " "
            if lv.find(_q("a:buAutoNum")) is not None:
                counters[level] = counters.get(level, 0) + 1
                return f"{counters[level]}. "
        return ""

    def _run_css(self, rpr, defaults, fallback, scale, default_color=None, default_bold=False) -> tuple[str, float]:
        specific = ([rpr] if rpr is not None else []) + defaults
        sources = specific + fallback
        size = int(_first_attr(sources, "sz", 0)) / 100 or DEFAULT_SIZE_PT
        size = max(size * scale, 1.0)
        bold = _first_attr(sources, "b", "1" if default_bold else "0") in ("1", "true")
        italic = _first_attr(sources, "i", "0") in ("1", "true")
        underline = _first_attr(sources, "u", "none") != "none"
        color = self._first_color(specific) or default_color or self._first_color(fallback)
        if not color:
            color = tuple(c / 255 for c in self.scheme.get(self.clr_map.get("tx1", "dk1"), (0, 0, 0)))
        face = ""
        for src in sources:
            latin = src.find(_q("a:latin"))
            if latin is not None and latin.get("typeface"):
                face = latin.get("typeface")
                break
        face = {"+mj-lt": self.fonts["major"], "+mn-lt": self.fonts["minor"]}.get(face, face)
        css = [f"font-size:{size:.1f}pt", f"color:#{''.join(f'{round(c * 255):02x}' for c in color)}",
               f"font-family:{_generic_family(face)}"]
        if bold:
            css.append("font-weight:bold")
        if italic:
            css.append("font-style:italic")
        if underline:
            css.append("text-decoration:underline")
        return ";".join(css), size


# ---------------------------------------------------------------- small helpers
def _hex(val) -> tuple[int, int, int] | None:
    try:
        return int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
    except (TypeError, ValueError, IndexError):
        return None


def _apply_mods(rgb, el) -> tuple[float, float, float]:
    """DrawingML colour tweaks (lumMod/lumOff for theme tints, shade, tint) as 0-1 floats."""
    r, g, b = (c / 255 for c in rgb)
    for m in el:
        name, val = _local(m.tag), int(m.get("val", 100000)) / 100000
        if name in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * val if name == "lumMod" else l + val
            r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0.0), 1.0), s)
        elif name == "shade":
            r, g, b = r * val, g * val, b * val
        elif name == "tint":
            r, g, b = (c + (1 - c) * (1 - val) for c in (r, g, b))
    return tuple(min(max(c, 0.0), 1.0) for c in (r, g, b))


def _tint(rgb, amount: float) -> tuple[float, float, float]:
    """`rgb` mixed with white: amount 0.2 keeps 20% of the colour."""
    return tuple(1 - (1 - c) * amount for c in rgb)


def _alpha(el) -> float:
    a = el.find(_q("a:alpha")) if el is not None else None
    return int(a.get("val", 100000)) / 100000 if a is not None else 1.0


def _first_attr(elements, name, default):
    for el in elements:
        if el is not None and el.get(name) is not None:
            return el.get(name)
    return default


def _anchor_of(el) -> str | None:
    if el is None:
        return None
    return el.get("anchor")


def _generic_family(face: str) -> str:
    low = face.lower()
    if any(k in low for k in ("mono", "courier", "consolas", "menlo", "code")):
        return "monospace"
    if any(k in low for k in ("times", "georgia", "garamond", "serif", "cambria", "palatino", "book")) \
            and "sans" not in low:
        return "serif"
    return "sans-serif"


def _crop_image(data: bytes, crop) -> bytes:
    """Apply PowerPoint's a:srcRect crop (percent x 1000 from each edge)."""
    try:
        im = Image.open(io.BytesIO(data))
        w, h = im.size
        l, t, r, b = (int(crop.get(k, 0)) / 100000 for k in ("l", "t", "r", "b"))
        box = (round(w * l), round(h * t), round(w * (1 - r)), round(h * (1 - b)))
        if box[2] <= box[0] or box[3] <= box[1]:
            return data
        buf = io.BytesIO()
        im.crop(box).save(buf, "PNG")
        return buf.getvalue()
    except Exception:
        return data


# ---------------------------------------------------------------- PDF -> PPTX
def pdf_to_pptx(src: str | Path) -> list[Path]:
    """One slide per page: the page as a full-slide picture, its text in the speaker notes
    (so it stays searchable and copyable)."""
    from pptx import Presentation

    src = Path(src)
    prs = Presentation()
    blank = prs.slide_layouts[6]
    with open_pdf(src) as doc:
        if doc.page_count == 0:
            raise UserError("This PDF has no pages.")
        if doc.page_count > MAX_PDF_PAGES:
            raise UserError(f"This PDF has {doc.page_count} pages; the limit for PowerPoint is {MAX_PDF_PAGES}. "
                            "Split it first.")
        first = doc[0].rect
        sw = min(max(round(first.width * EMU_PER_PT), PPTX_MIN_EMU), PPTX_MAX_EMU)
        sh = min(max(round(first.height * EMU_PER_PT), PPTX_MIN_EMU), PPTX_MAX_EMU)
        prs.slide_width, prs.slide_height = sw, sh
        for page in doc:
            slide = prs.slides.add_slide(blank)
            png = page.get_pixmap(dpi=150, alpha=False).tobytes("png")
            # pages can differ in size: fit each one inside the slide, centred
            k = min(sw / (page.rect.width * EMU_PER_PT), sh / (page.rect.height * EMU_PER_PT))
            w, h = round(page.rect.width * EMU_PER_PT * k), round(page.rect.height * EMU_PER_PT * k)
            slide.shapes.add_picture(io.BytesIO(png), (sw - w) // 2, (sh - h) // 2, w, h)
            text = CONTROL_CHARS.sub("", page.get_text("text")).strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text
    out = sibling(src, ext=".pptx")
    atomic_save(prs.save, out)
    return [out]
