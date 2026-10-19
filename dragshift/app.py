"""DragShift: Shift+drag any file to convert or edit it from a radial liquid-glass menu."""
from __future__ import annotations

import json
import os
import sys
import threading
import time
import traceback
from pathlib import Path

from PySide6.QtCore import QDir, QObject, QPointF, QRunnable, Qt, QThreadPool, QTimer, QUrl, Signal
from PySide6.QtGui import (QAction, QColor, QCursor, QDesktopServices, QGuiApplication, QIcon, QPainter,
                           QPixmap)
from PySide6.QtNetwork import QLocalServer, QLocalSocket
from PySide6.QtWidgets import QApplication, QFileDialog, QMenu, QSystemTrayIcon, QWidget

from . import __version__, actions, archives, compress, convert, glass, imaging, media, media_windows, qr, report, tools
from .reorder import ReorderWindow
from .floating import Floats
from .radial import Item, RadialMenu
from .settings import Settings, SettingsWindow, WelcomeWindow, key_glyphs
from .toasts import ToastStack, trash, trash_message
from .trigger import DragShiftWatcher
from .util import UserError, category, file_size, human_size, keep, plural, release

APP_NAME = "DragShift"
TIP_URL = "https://cash.app/$AtikaKadesha"
VERSION = __version__
IS_WIN = sys.platform == "win32"
IS_MAC = sys.platform == "darwin"
from .handoff import ipc_name
IPC_NAME = ipc_name()

GLYPHS = {"Compress": "⇲", "Metadata": "ⓘ", "Redact": "▮", "Crop": "⌗", "Edit": "☼",
          "Annotate": "✎", "Add BG": "▣", "Resize": "⤢", "Split": "✂", "Rotate": "⟳",
          "Merge": "⧉", "Protect": "⚿", "Extract": "⧈", "Audio": "♪", "Mute": "⊘", "Frame": "◧",
          "Normalize": "≈", "Strip Tags": "⌫", "Trim": "⇹", "Speed": "»",
          "Read QR": "▦", "Reorder": "⇅", "Extract Files": "⇱",
          "AirDrop": "◉", "Float": "⌖", "Copy": "⧉", "Move": "➜", "Copy Path": "⎘", "Terminal": "›_",
          "Zip": "▤"}

HINTS = {
    "Compress": "Shrink file size", "Metadata": "View · remove GPS", "Redact": "Hide private info",
    "Crop": "Trim to a ratio", "Edit": "Light · colour · rotate", "Annotate": "Arrows · text · steps",
    "Add BG": "Gradient backdrop", "Resize": "Change dimensions", "Split": "One PDF per page",
    "Rotate": "Turn pages 90°", "Merge": "Join into one PDF", "Protect": "Add a password",
    "Extract": "Save embedded images", "Audio": "Save the sound", "Mute": "Remove the sound",
    "Frame": "Save a still as PNG", "Normalize": "Even out loudness", "Strip Tags": "Remove tags · cover art",
    "Trim": "Cut the start or end", "Speed": "Faster or slower",
    "Read QR": "Show what a code says", "Reorder": "Drag pages into order", "Extract Files": "Into a new folder",
    "AirDrop": "Send to a nearby device", "Float": "Keep it handy to drag later", "Copy": "Copy to a folder…",
    "Move": "Move to a folder…", "Copy Path": "Path to the clipboard", "Terminal": "Open a shell here",
    "Zip": "Pack it all in a .zip",
}
# same tool name, different meaning for this kind of file
CATEGORY_HINTS = {"video": {"Rotate": "Turn 90° clockwise", "Crop": "Cut away the edges"},
                  "audio": {"Compress": "Lower the bitrate"}}
FORMAT_HINTS = {"JPG": "Small, universal", "PNG": "Lossless · transparency", "WEBP": "Smallest for web",
                "HEIC": "iPhone format", "TIFF": "Print · archival", "AVIF": "Next-gen, tiny",
                "BMP": "Uncompressed", "PDF": "Document", "ICO": "Windows icon", "GIF": "Legacy web",
                "DOCX": "Editable Word", "TXT": "Plain text",
                "MP4": "Plays everywhere", "MOV": "Apple · editing", "MKV": "Keeps every track",
                "WEBM": "Web video", "AVI": "Older players", "MP3": "Plays everywhere",
                "M4A": "Apple · small", "WAV": "Uncompressed", "FLAC": "Lossless, smaller",
                "OGG": "Open format", "OPUS": "Small, great quality", "AIFF": "Uncompressed · Apple",
                "XLSX": "Excel workbook", "ODS": "OpenDocument sheet", "CSV": "Comma-separated",
                "TSV": "Tab-separated", "PPTX": "PowerPoint", "EPUB": "Ebook", "MD": "Markdown",
                "HTML": "Web page", "ZIP": "Opens everywhere", "TAR": "Unix · keeps permissions",
                "TAR.GZ": "Smaller tar", "SRT": "Most players", "VTT": "Web video"}
CATEGORY_FORMAT_HINTS = {"video": {"GIF": "Silent loop · 640 px", "MP3": "Sound only", "M4A": "Sound only · small",
                                   "WAV": "Sound only · lossless"},
                         "pdf": {"XLSX": "Tables to sheets", "PPTX": "One slide per page", "EPUB": "Reflowed ebook",
                                 "MD": "Text with headings"},
                         "sheet": {"CSV": "One file per sheet", "TSV": "One file per sheet", "PDF": "Printable table",
                                   "DOCX": "Word table"},
                         "slides": {"PDF": "One page per slide", "JPG": "One image per slide",
                                    "PNG": "One image per slide"},
                         "pages": {"PDF": "Printable copy"},
                         "subtitle": {"TXT": "Dialogue only"},
                         "svg": {"PDF": "Vector · sharp at any zoom"}}

IMAGE_TOOLS = ["Compress", "Resize", "Crop", "Edit", "Annotate", "Redact", "Add BG", "Metadata", "Read QR"]
PDF_TOOLS = ["Compress", "Split", "Rotate", "Protect", "Extract", "Metadata", "Reorder"]
ARCHIVE_TOOLS = ["Extract Files"]
# tool name -> (function run on each file, spinner text), or the tool window that opens instead
MEDIA_TOOLS = {
    "video": {"Compress": compress.video_window, "Trim": media_windows.TrimWindow,
              "Crop": media_windows.VideoCropWindow, "Speed": media_windows.SpeedWindow,
              "Audio": (media.extract_audio, "Extracting audio…"), "Mute": (media.mute, "Removing sound…"),
              "Frame": media_windows.FrameWindow, "Rotate": (media.rotate, "Rotating…")},
    "audio": {"Compress": compress.audio_window, "Trim": media_windows.TrimWindow,
              "Speed": media_windows.SpeedWindow, "Normalize": (media.normalize, "Evening out volume…"),
              "Strip Tags": (media.strip_metadata, "Removing tags…")},
}


# ---------------------------------------------------------------- background jobs
class _Sig(QObject):
    done = Signal(object)
    failed = Signal(str)
    progress = Signal(float)


class Job(QRunnable):
    def __init__(self, action: str, srcs: list, fn, extra: str = ""):
        super().__init__()
        self.action, self.srcs, self.fn, self.extra = action, srcs, fn, extra
        self.sig = _Sig()
        self.cancel = threading.Event()

    def run(self):
        try:
            with media.reporting(self.sig.progress.emit, self.cancel):
                outs = self.fn() or []
            self.sig.done.emit(report.summarize(self.action, self.srcs, outs, self.extra))
        except UserError as e:
            self.sig.failed.emit(str(e))
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            self.sig.failed.emit(friendly_error(e))


def friendly_error(e: Exception) -> str:
    msg = str(e) or e.__class__.__name__
    low = msg.lower()
    if isinstance(e, PermissionError) or "permission denied" in low:
        return "No permission to write next to this file. Try a folder you own, or set Settings → Save results → subfolder."
    if "cannot identify image" in low:
        return "This file doesn't look like an image DragShift can read."
    if isinstance(e, FileNotFoundError):
        return "The file was moved or deleted."
    return msg


def tool_items(names: list[str], cat: str, kind: str = "tool") -> list[Item]:
    hints = {**HINTS, **CATEGORY_HINTS.get(cat, {})}
    return [Item(n, GLYPHS[n], key=f"{kind}:{n}", hint=hints[n]) for n in names]


def file_category(path: str) -> str | None:
    """Category for Convert/Tools; folders have none (even one called "Photos.png")."""
    return None if os.path.isdir(path) else category(path)


def archive_items(paths: list[str]) -> list[Item]:
    """Archive formats for a selection nothing else converts (folders, unknown files): one
    archive of everything."""
    n = len(paths)
    return [Item(t, key=f"archive:{t}", hint=f"{n} items → 1 {t}" if n > 1 else FORMAT_HINTS[t])
            for t in archives.TARGETS]


def primary_category(paths: list[str]) -> str | None:
    """Convert/Tools follow the first item they can handle; the rest of the selection only
    gets the Actions ring."""
    return next((c for c in map(file_category, paths) if c is not None), None)


# ---------------------------------------------------------------- icon
def make_tray_icon() -> QIcon:
    """macOS menu bar wants a monochrome 'template' image; Windows gets the colour icon."""
    if not IS_MAC:
        return make_icon()
    pm = QPixmap(44, 44)
    pm.fill(Qt.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing)
    from PySide6.QtGui import QPen
    p.setPen(QPen(QColor(0, 0, 0), 4.5))
    p.drawEllipse(QPointF(22, 22), 14, 14)
    p.setPen(Qt.NoPen)
    p.setBrush(QColor(0, 0, 0))
    p.drawEllipse(QPointF(22, 22), 5, 5)
    for a in range(0, 360, 45):  # petal gaps
        import math
        x, y = 22 + 14 * math.cos(math.radians(a + 22.5)), 22 + 14 * math.sin(math.radians(a + 22.5))
        p.setCompositionMode(QPainter.CompositionMode_Clear)
        p.drawEllipse(QPointF(x, y), 2.2, 2.2)
        p.setCompositionMode(QPainter.CompositionMode_SourceOver)
    p.end()
    pm.setDevicePixelRatio(2.0)
    icon = QIcon(pm)
    icon.setIsMask(True)
    return icon


def make_icon() -> QIcon:
    ico = Path(__file__).with_name("assets") / "dragshift.ico"
    if ico.exists():
        return QIcon(str(ico))
    pm = QPixmap(64, 64)
    pm.fill(Qt.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(glass.iridescent_pen(QPointF(32, 32), 7, 255))
    p.drawEllipse(QPointF(32, 32), 24, 24)
    p.end()
    return QIcon(pm)


# ---------------------------------------------------------------- the app
class DragShiftApp(QObject):
    def __init__(self, qapp: QApplication):
        super().__init__()
        self.qapp = qapp
        self.settings = Settings()
        self.settings.apply()
        self.pool = QThreadPool.globalInstance()
        self.windows: list[QWidget] = []
        self.toasts = ToastStack()
        self.floats = Floats()
        self.menu = RadialMenu(self.build_rings, self.describe)
        self.menu.activated.connect(self.run_item)

        # the watcher only manages menus it opened (drag mode); click-mode menus (tray, right-click,
        # Send To) must not be closed just because no mouse button is held
        self.watcher = DragShiftWatcher(lambda: self.menu.isVisible() and not self.menu.click_mode,
                                        lambda: self.menu.got_drag, self)
        self.apply_keys()
        self.watcher.show_menu.connect(
            lambda pos, ring: self.menu.popup(pos, ring=ring, dragged=self.watcher.drag_paths()))
        self.watcher.switch_ring.connect(self.menu.ring_key)
        self.watcher.nav_key.connect(self.menu.nav)
        self.watcher.hide_menu.connect(self.menu.close_menu)
        self.watcher.start()

        # files handed over by "Send to", the right-click menu, or a second launch
        self._pending: list[str] = []
        self._pending_timer = QTimer(self, singleShot=True, interval=280, timeout=self._flush_pending)
        self.server = listen_ipc(IPC_NAME, self)
        self.server.newConnection.connect(self._ipc)

        self.tray = QSystemTrayIcon(make_tray_icon(), self)
        self.tray.setToolTip(f"{APP_NAME} — press {self.trigger_label()} while dragging a file")
        m = QMenu()
        m.setStyleSheet(glass.QSS)
        self.act_enabled = QAction("Enabled", m, checkable=True, checked=True)
        self.act_enabled.toggled.connect(lambda on: setattr(self.watcher, "enabled", on))
        m.addAction(self.act_enabled)
        m.addAction("Open a file…", self.open_dialog)
        m.addSeparator()
        m.addAction("Settings…", self.open_settings)
        m.addAction("How to use", self.show_welcome)
        m.addAction("Tip the developer ♥", lambda: QDesktopServices.openUrl(QUrl(TIP_URL)))
        m.addSeparator()
        m.addAction(f"Quit {APP_NAME}", self.qapp.quit)
        self.tray_menu = m
        self.tray.setContextMenu(m)
        self.tray.activated.connect(self._tray_click)
        self.tray.show()

    def trigger_label(self) -> str:
        return self.settings.key_label("Convert")

    def apply_keys(self):
        """Hand the wheel -> key map to the watcher, and key glyphs to the hub hint."""
        keys = self.settings.key_map()
        self.watcher.set_keys(keys)
        glyphs = key_glyphs()
        self.menu.ring_keys = {wheel: glyphs.get(key, "") for wheel, key in keys.items()}
        self.menu.wheel_keys = dict(keys)

    # ------------------------------------------------------------ IPC
    def _ipc(self):
        sock = self.server.nextPendingConnection()
        if sock is None:
            return
        buf = bytearray()
        done = {"v": False}

        def consume():
            if done["v"]:
                return
            try:
                buf.extend(bytes(sock.readAll()))
            except RuntimeError:
                pass
            try:
                data = json.loads(buf.decode("utf-8"))
            except Exception:
                return  # incomplete: wait for more
            done["v"] = True
            request = ipc_request(data)
            files, show = request or ([], False)
            print("IPC: ignored a message that isn't from DragShift" if request is None
                  else f"IPC: {len(files)} file(s) handed over", flush=True)
            if files:
                self.queue_files(files)
            elif show:
                self.toasts.info("DragShift is already running",
                                 f"Drag any file and press {self.trigger_label()}. It lives in the tray.")
            sock.disconnectFromServer()
            sock.deleteLater()

        # Windows pipes: the sender may already be gone, so read on every signal + a fallback timer
        sock.readyRead.connect(consume)
        sock.disconnected.connect(consume)
        for ms in (0, 50, 200, 600):
            QTimer.singleShot(ms, consume)

    def queue_files(self, files: list[str]):
        """Explorer's right-click menu launches once per file: gather them, then show one menu."""
        for f in files:
            if f not in self._pending:
                self._pending.append(f)
        self._pending_timer.start()

    def _flush_pending(self):
        files, self._pending = self._pending, []
        if files:
            self.show_click_menu(files)

    # ------------------------------------------------------------ menu contents
    def describe(self, paths: list[str]) -> str:
        if len(paths) > 1:
            noun = "items" if any(os.path.isdir(p) for p in paths) else "files"
            return f"{len(paths)} {noun}"
        if os.path.isdir(paths[0]):
            return "Folder"
        return human_size(file_size(paths[0]))

    def build_rings(self, paths: list[str]) -> dict[str, list[Item]]:
        """Convert and Tools when the first file DragShift can handle has them; Actions for
        anything, folders and unsupported files included."""
        actions_ring = {"Actions": tool_items(actions.available(), "", kind="action")}
        cat = primary_category(paths)
        if cat is None:
            return {"Convert": archive_items(paths), **actions_ring}
        same = [p for p in paths if file_category(p) == cat]
        ext = archives.archive_suffix(same[0]) if cat == "archive" else Path(same[0]).suffix
        convert_items = []
        format_hints = {**FORMAT_HINTS, **CATEGORY_FORMAT_HINTS.get(cat, {})}
        for t in convert.targets_for(cat, ext):
            hint = format_hints.get(t, "")
            if cat == "image" and t == "PDF" and len(same) > 1:
                hint = f"{len(same)} images → 1 PDF"
            convert_items.append(Item(t, key=f"convert:{t}", hint=hint))
        rings: dict[str, list[Item]] = {"Convert": convert_items}
        if cat == "image":
            rings["Tools"] = tool_items(IMAGE_TOOLS, cat)
        elif cat == "pdf":
            rings["Tools"] = tool_items(PDF_TOOLS[:3] + (["Merge"] if len(same) > 1 else []) + PDF_TOOLS[3:], cat)
        elif cat in MEDIA_TOOLS:
            rings["Tools"] = tool_items(list(MEDIA_TOOLS[cat]), cat)
        elif cat == "archive":
            rings["Tools"] = tool_items(ARCHIVE_TOOLS, cat)
        return {**rings, **actions_ring}

    # ------------------------------------------------------------ actions
    def run_item(self, item: Item, paths: list[str]):
        kind, name = item.key.split(":", 1)
        if kind == "action":
            self.run_action(name, paths)
            return
        if kind == "archive":       # the whole selection, folders included
            self.make_archive(name, paths)
            return
        cat = primary_category(paths)
        paths = [p for p in paths if file_category(p) == cat]
        if kind == "convert":
            if cat == "image" and name == "PDF" and len(paths) > 1:
                self.submit(f"convert:{name}", paths, lambda: convert.images_to_single_pdf(paths),
                            f"Combining {len(paths)} images…")
            else:
                verb = f"Converting to {name}…"
                self.submit(f"convert:{name}", paths,
                            lambda: [o for p in paths for o in convert.convert(p, cat, name)], verb)
            return
        if cat in MEDIA_TOOLS:
            tool = MEDIA_TOOLS[cat][name]
            if isinstance(tool, tuple):
                fn, working = tool
                self.submit(f"tool:{name}", paths, lambda: [fn(p) for p in paths], working)
            elif name == "Compress":   # the one media window that takes the whole selection
                self.open_tool(name, lambda: tool(paths, self.notify))
            else:
                self.open_tool(name, lambda: tool(paths[0], self.notify))
            return
        if cat == "archive":
            self.submit("tool:Extract Files", paths, lambda: [archives.extract(p) for p in paths], "Extracting…")
            return
        if cat == "pdf":
            if name == "Merge":
                self.submit("tool:Merge", paths, lambda: [imaging.merge_pdfs(paths)], "Merging…")
            elif name == "Metadata":
                self.open_window(tools.MetadataWindow(paths[0], self.notify, is_pdf=True))
            elif name == "Protect":
                self.open_window(tools.ProtectWindow(paths[0], self.notify))
            elif name == "Compress":
                self.open_tool(name, lambda: compress.pdf_window(paths, self.notify))
            elif name == "Reorder":
                self.open_tool(name, lambda: ReorderWindow(paths[0], self.notify))
            else:
                fns = {"Split": imaging.split_pdf, "Rotate": lambda p: [imaging.rotate_pdf(p)],
                       "Extract": imaging.extract_pdf_images}
                f = fns[name]
                self.submit(f"tool:{name}", paths, lambda: [o for p in paths for o in f(p)],
                            {"Split": "Splitting…", "Rotate": "Rotating…", "Extract": "Extracting images…"}[name])
            return
        windows = {"Crop": tools.CropWindow, "Add BG": tools.AddBgWindow, "Annotate": tools.AnnotateWindow,
                   "Redact": tools.RedactWindow, "Edit": tools.EditWindow, "Metadata": tools.MetadataWindow,
                   "Read QR": qr.QrWindow}
        if name == "Compress":
            self.open_tool(name, lambda: tools.CompressWindow(paths, self.notify))
        elif name == "Resize":
            self.open_tool(name, lambda: tools.ResizeWindow(paths, self.notify))
        else:
            self.open_tool(name, lambda: windows[name](paths[0], self.notify))

    def open_tool(self, name: str, make):
        """Build a tool window and show it; a file it can't read becomes an error card."""
        try:
            self.open_window(make())
        except UserError as e:
            self.toasts.error(f"Couldn't open {name}", str(e))
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            self.toasts.error(f"Couldn't open {name}", friendly_error(e))

    def run_action(self, name: str, paths: list[str]):
        """Actions work on every dragged item: files of any type, folders, mixed selections."""
        try:
            if name == "AirDrop":
                actions.airdrop(paths)
            elif name == "Float":
                self.floats.open(paths)
            elif name in ("Copy", "Move"):
                self.copy_or_move(name, paths)
            elif name == "Copy Path":
                actions.copy_paths(paths)
                self.toasts.info("Copied path" if len(paths) == 1 else f"Copied {len(paths)} paths",
                                 actions.path_text(paths), 2500)
            elif name == "Terminal":
                actions.open_terminal(paths)
            elif name == "Zip":
                self.make_archive("ZIP", paths)
        except UserError as e:
            self.toasts.error(f"Couldn't {actions.VERBS[name]}", str(e))
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            self.toasts.error(f"Couldn't {actions.VERBS[name]}", friendly_error(e))

    def make_archive(self, fmt: str, paths: list[str]):
        """One archive of every dragged item (files and folders), beside the first one."""
        self.submit(f"archive:{fmt}", paths, lambda: [archives.create(paths, fmt)], "Compressing…")

    def copy_or_move(self, name: str, paths: list[str]):
        """Ask for a folder (starting at the last one used), then copy or move in the background.
        Copy's Undo moves the copies to the Trash; Move's Undo moves the items back."""
        if IS_MAC:
            from . import platform_mac
            platform_mac.activate()   # a menu-bar app's dialog would open behind other windows
        last = self.settings.get("last_folder")
        what = Path(paths[0]).name if len(paths) == 1 else self.describe(paths)
        dest = QFileDialog.getExistingDirectory(None, f"{name} {what} to",
                                                last if last and os.path.isdir(last) else QDir.homePath())
        if not dest:
            return
        self.settings.set("last_folder", dest)
        folder = Path(dest).name or dest
        if name == "Copy":
            copied: list[Path] = []

            def copy() -> list[Path]:
                copied.extend(actions.copy_to(paths, dest))
                return copied

            # prune=False: never remove the user's destination folder, even if it ends up empty
            self.submit("action:Copy", paths, copy, "Copying…", extra=folder,
                        undo=lambda: trash_message(trash(copied, prune=False)))
        else:
            moves = actions.Moves()

            def undo_move() -> str:
                n = actions.undo_move(moves)
                return f"Moved {plural(n, 'item')} back"

            self.submit("action:Move", paths, lambda: actions.move_to(paths, dest, moves), "Moving…",
                        extra=folder, undo=undo_move)

    def open_window(self, w: QWidget):
        if hasattr(w, "submit"):
            w.submit = self.submit
        self.windows.append(w)
        w.destroyed.connect(lambda *_: self.windows.remove(w) if w in self.windows else None)
        if hasattr(w, "present"):
            w.present()
        else:
            w.show()

    def submit(self, action: str, srcs: list, fn, working: str = "Working…", extra: str = "", undo=None):
        """Run in the background. A spinner card appears only if it takes > 150 ms; jobs that
        report progress (ffmpeg) show a percentage, and closing the card cancels them.
        `undo` replaces the result card's default Undo (moving the outputs to the Trash)."""
        state = {"card": None, "finished": False, "cancelled": False}
        names = Path(srcs[0]).name if len(srcs) == 1 else f"{len(srcs)} files"
        job = Job(action, srcs, fn, extra)

        def cancel():
            state["cancelled"] = True
            job.cancel.set()

        def forget(_card):  # the card was closed or pushed off the stack: don't touch it again
            state["card"] = None

        def show_spinner():
            if not state["finished"]:
                card = state["card"] = self.toasts.new()
                card.working(working, names)
                card.cancel_requested.connect(cancel)
                card.closed.connect(forget)

        def progressed(frac):
            if state["card"] is not None and not state["finished"]:
                state["card"].working(f"{working} {frac:.0%}", names)

        QTimer.singleShot(150, show_spinner)

        def done(r):
            release(job)
            state["finished"] = True
            if undo is not None:
                r.undo = undo
            card = state["card"] or self.toasts.new()
            card.done(r)

        def failed(msg):
            release(job)
            state["finished"] = True
            if state["cancelled"]:
                self.toasts.info("Cancelled", names, 2500)
                return
            card = state["card"] or self.toasts.new()
            card.error("Couldn't finish", msg)

        job.sig.done.connect(done)
        job.sig.failed.connect(failed)
        job.sig.progress.connect(progressed)
        self.pool.start(keep(job))

    def notify(self, action: str, srcs: list, outs: list, extra: str = ""):
        """Called by tool windows after they save."""
        self.toasts.new().done(report.summarize(action, srcs, outs, extra))

    # ------------------------------------------------------------ misc
    def _tray_click(self, reason):
        # macOS: a click on the menu-bar icon shows the menu; Windows: click opens a file
        if not IS_MAC and reason in (QSystemTrayIcon.DoubleClick, QSystemTrayIcon.Trigger):
            self.open_dialog()

    def open_dialog(self):
        files, _ = QFileDialog.getOpenFileNames(None, "Choose file(s) for DragShift", QDir.homePath())
        if files:
            self.show_click_menu(files)

    def show_click_menu(self, files: list[str]):
        g = (QGuiApplication.screenAt(QCursor.pos()) or QGuiApplication.primaryScreen()).availableGeometry()
        pos = QCursor.pos() if g.contains(QCursor.pos()) else g.center()
        self.menu.popup(pos, files)
        self.menu.activateWindow()

    def open_settings(self):
        w = SettingsWindow(self.settings)
        w.changed.connect(self._settings_changed)
        self.open_window(w)

    def _settings_changed(self):
        self.apply_keys()
        self.tray.setToolTip(f"{APP_NAME} — press {self.trigger_label()} while dragging a file")

    def show_welcome(self):
        self.open_window(WelcomeWindow(self.settings))


class DragShiftQApplication(QApplication):
    """macOS delivers files from Finder's "Open With", the Dock and `open -a` as FileOpen
    events (not argv)."""
    file_opened = Signal(str)

    def __init__(self, argv):
        super().__init__(argv)
        self.early_files: list[str] = []

    def event(self, e):
        from PySide6.QtCore import QEvent
        if e.type() == QEvent.FileOpen:
            f = e.file()
            if f:
                self.early_files.append(f)
                self.file_opened.emit(f)
            return True
        return super().event(e)


def mac_setup(qapp):
    from PySide6.QtGui import QFont
    for fam in ("Segoe UI Variable Display", "Segoe UI Variable Text", "Segoe UI"):
        QFont.insertSubstitutions(fam, [".AppleSystemUIFont", "SF Pro Display", "Helvetica Neue"])
    try:
        from . import platform_mac
        platform_mac.become_accessory()
    except Exception:
        pass


def listen_ipc(name: str, parent=None) -> QLocalServer:
    """The server a second launch hands its files to. On macOS and Linux it's a socket file, and
    UserAccessOption (set before listen) makes it this user's only: no other account on the
    machine can connect and send files."""
    server = QLocalServer(parent)
    if not IS_WIN:
        server.setSocketOptions(QLocalServer.UserAccessOption)
    QLocalServer.removeServer(name)
    server.listen(name)
    return server


def ipc_request(data) -> tuple[list[str], bool] | None:
    """(files that exist, show) from a hand-off message ({"files": [...]} or {"show": true}),
    None for anything else: the socket takes whatever a local program sends."""
    if not isinstance(data, dict):
        return None
    files, show = data.get("files", []), data.get("show", False)
    if not isinstance(files, list) or not all(isinstance(f, str) for f in files) or not isinstance(show, bool):
        return None
    return [f for f in files if os.path.exists(f)], show


def send_to_running(payload: dict) -> bool:
    s = QLocalSocket()
    s.connectToServer(IPC_NAME)
    if not s.waitForConnected(300):
        return False
    s.write(json.dumps(payload).encode("utf-8"))
    s.flush()
    s.waitForBytesWritten(500)
    s.disconnectFromServer()
    return True


def setup_logging():
    """Errors go to %LOCALAPPDATA%\\DragShift\\log.txt (pythonw / the .exe have no console)."""
    import faulthandler
    if IS_MAC:
        base = Path.home() / "Library" / "Logs" / "DragShift"
    else:
        base = Path(os.environ.get("LOCALAPPDATA") or Path.home() / ".local" / "share") / "DragShift"
    base.mkdir(parents=True, exist_ok=True)
    path = base / "log.txt"
    if path.exists() and path.stat().st_size > 1_000_000:
        path.replace(base / "log.old.txt")
    f = open(path, "a", encoding="utf-8", buffering=1)
    f.write(f"\n==== DragShift {VERSION} started {time.strftime('%Y-%m-%d %H:%M:%S')} ====\n")
    if sys.stderr is None or getattr(sys, "frozen", False) or sys.executable.lower().endswith("pythonw.exe"):
        sys.stdout = sys.stderr = f
    faulthandler.enable(f)

    def hook(t, v, tb):
        f.write("Unhandled: " + "".join(traceback.format_exception(t, v, tb)))
    sys.excepthook = hook
    return path


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv if argv is None else argv
    try:
        setup_logging()
    except Exception:
        pass
    if IS_WIN:
        try:  # crisp taskbar/tray icon grouping
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("DragShift.App")
        except Exception:
            pass
    QApplication.setHighDpiScaleFactorRoundingPolicy(Qt.HighDpiScaleFactorRoundingPolicy.PassThrough)
    qapp = DragShiftQApplication(argv)
    if IS_MAC:
        mac_setup(qapp)
    qapp.setApplicationName(APP_NAME)
    qapp.setApplicationVersion(VERSION)
    qapp.setQuitOnLastWindowClosed(False)
    qapp.setWindowIcon(make_icon())

    if "--selftest" in argv:
        return selftest()
    files = [os.path.abspath(a) for a in argv[1:] if not a.startswith("-") and os.path.exists(a)]
    # already running? hand over and exit (instant: no second UI start-up)
    if send_to_running({"files": files} if files else {"show": "--quiet" not in argv}):
        return 0
    app = DragShiftApp(qapp)
    qapp.file_opened.connect(lambda f: app.queue_files([f]))
    if IS_MAC:  # Finder > Services > "Open in DragShift", or the shortcut the user gave it
        from . import platform_mac
        platform_mac.register_services(app.queue_files)
    if qapp.early_files:
        app.queue_files(list(qapp.early_files))
    if files:  # e.g. "Send to → DragShift" or drag onto the .exe
        QTimer.singleShot(120, lambda: app.queue_files(files))
    elif not app.settings.get("welcomed"):
        QTimer.singleShot(300, app.show_welcome)
    elif "--quiet" not in argv:
        QTimer.singleShot(400, lambda: app.toasts.info(
            "DragShift is running", f"Drag any file, then press {app.trigger_label()}."))
    return qapp.exec()


def selftest_documents(d: Path, docx: Path):
    """Office, ebook and text formats: one pass through every reader and writer library."""
    (d / "s.csv").write_text("name,qty\nZürich,3\n", encoding="utf-8")
    for t in convert.targets_for("sheet", ".csv"):
        convert.convert(d / "s.csv", "sheet", t)
    (d / "m.md").write_text("# Title\n\n- one\n- two\n", encoding="utf-8")
    epub = convert.convert(d / "m.md", "markdown", "EPUB")[0]
    for t in convert.targets_for("epub", ".epub"):
        convert.convert(epub, "epub", t)
    for t in ("MD", "HTML"):
        convert.convert(docx, "docx", t)
    from pptx import Presentation
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Hello"
    deck.save(d / "deck.pptx")
    for t in convert.targets_for("slides", ".pptx"):
        convert.convert(d / "deck.pptx", "slides", t)
    pdf = convert.convert(d / "deck.pptx", "slides", "PDF")[0]
    for t in convert.PDF_TARGETS:
        convert.convert(pdf, "pdf", t)


def selftest_extras(d: Path):
    """SVG (QtSvg plugin), QR (zxing-cpp native module), archives and subtitles."""
    from . import archives, qr, subtitles
    (d / "v.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="40" height="30">'
                             '<rect width="40" height="30" fill="red"/></svg>', encoding="utf-8")
    for t in convert.targets_for("svg", ".svg"):
        convert.convert(d / "v.svg", "svg", t)
    import numpy as np
    import zxingcpp
    from PIL import Image
    code = zxingcpp.write_barcode_to_image(zxingcpp.create_barcode("DragShift", zxingcpp.BarcodeFormat.QRCode), scale=4)
    Image.fromarray(np.asarray(code)).save(d / "q.png")
    if [c.text for c in qr.decode(d / "q.png")] != ["DragShift"]:
        raise RuntimeError("QR decode returned the wrong text")
    zipped = archives.create([d / "v.svg", d / "q.png"], "ZIP")
    archives.convert(zipped, "TAR.GZ")
    if not (archives.extract(zipped) / "q.png").exists():
        raise RuntimeError("extracted archive is missing a file")
    (d / "s.srt").write_text("1\n00:00:01,000 --> 00:00:02,500\nHello\n", encoding="utf-8")
    for t in subtitles.targets_for(".srt"):
        subtitles.convert(d / "s.srt", t)


def selftest() -> int:
    """Used by CI on the frozen .exe: exercise every library and build the UI, then exit."""
    import tempfile
    from PIL import Image
    try:
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "t.png"
            Image.new("RGBA", (64, 48), (255, 100, 80, 255)).save(src)
            for t in convert.targets_for("image", ".png"):
                convert.convert(src, "image", t)
            pdf = convert.convert(src, "image", "PDF")[0]
            for t in ("JPG", "PNG", "DOCX", "TXT"):  # a picture-only PDF; text targets run in selftest_documents
                convert.convert(pdf, "pdf", t)
            (Path(d) / "n.txt").write_text("hello")
            docx = convert.convert(Path(d) / "n.txt", "text", "DOCX")[0]
            convert.convert(docx, "docx", "TXT")
            selftest_documents(Path(d), docx)
            selftest_extras(Path(d))
            imaging.add_background(convert.open_image(src), imaging.Background())
            imaging.compress_pdf(pdf)
            imaging.protect_pdf(pdf, "pw")
            if media.available():
                clip = media.make_test_clip(Path(d) / "clip.mp4")
                for t in convert.targets_for("video", ".mp4"):
                    convert.convert(clip, "video", t)
                for tool in MEDIA_TOOLS["video"].values():
                    if isinstance(tool, tuple):
                        tool[0](clip)
                    else:
                        tool(str(clip), lambda *a, **k: None).close()
                print(f"ffmpeg OK: {media.find_tool('ffmpeg')}")
            else:
                print("ffmpeg not found: video and audio skipped")
            from . import liquid
            liquid.render(Image.new("RGB", (80, 80)), liquid.sdf_circle(80, 80, 40, 40, 30))
            menu = RadialMenu(lambda p: {"Convert": [Item("PNG")]}, lambda p: "1 KB")
            menu._load([str(src)])
            n = lambda *a, **k: None
            for w in (tools.CropWindow(str(src), n), tools.AddBgWindow(str(src), n), tools.EditWindow(str(src), n),
                      tools.AnnotateWindow(str(src), n), tools.RedactWindow(str(src), n),
                      tools.MetadataWindow(str(src), n), tools.CompressWindow([str(src)], n),
                      tools.ResizeWindow([str(src)], n), tools.ProtectWindow(str(pdf), n)):
                w.close()
            ToastStack().info("ok", "ok").dismiss()
        print("SELFTEST OK")
        return 0
    except Exception:
        traceback.print_exc()
        print("SELFTEST FAILED")
        return 1
