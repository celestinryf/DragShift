"""Office, ebook and text formats: spreadsheets, slides, EPUB, Markdown, HTML, RTF and Pages.

Text documents go through one in-between form, `Doc`: a title plus an HTML body fragment (and
the pictures it uses). Each source format has a reader that makes a Doc (DOCX, Markdown, plain
text, HTML, EPUB, RTF, PDF) and each target has a writer that takes one (PDF, DOCX, Markdown,
TXT, HTML, EPUB). Spreadsheets live in sheets.py and slides in slides.py; this module routes
conversions to them."""
from __future__ import annotations

import base64
import html
import io
import os
import posixpath
import re
import sys
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath, PureWindowsPath
from urllib.parse import unquote

from . import sheets
from .convert import DOCX_TARGETS, PDF_TARGETS, TEXT_TARGETS, atomic_save, fitz, heading_level, html_to_pdf
from .convert import open_pdf, pdf_page_items, read_text, save_bytes, save_image, save_text
from .util import CONTROL_CHARS, HTML_EXTS, UserError, plural, sibling

SHEET_TARGETS = ["XLSX", "ODS", "CSV", "TSV", "PDF", "DOCX"]
SLIDES_TARGETS = ["PDF", "JPG", "PNG"]
EPUB_TARGETS = ["PDF", "DOCX", "TXT", "MD"]
MARKDOWN_TARGETS = ["PDF", "DOCX", "HTML", "EPUB", "TXT"]
RTF_TARGETS = ["TXT", "DOCX", "PDF", "MD"]
PAGES_TARGETS = ["PDF"]
HTML_TARGETS = ["PDF", "DOCX", "MD", "TXT", "EPUB"]
PLAIN_TARGETS = ["PDF", "DOCX", "HTML", "EPUB"]   # .txt

MAX_ZIP_ENTRY = 200 * 1024 * 1024   # bytes read from one file inside an EPUB/Pages zip (zip bombs)
MAX_IMAGE_BYTES = 50 * 1024 * 1024  # a single picture pulled into a document
# All chapters and pictures of one EPUB together. The whole book is held in memory as one HTML
# string and parsed again by each writer, so this bounds memory; real books, even picture-heavy
# ones, stay well under it, while a small zip can claim far more.
MAX_BOOK_BYTES = 300 * 1024 * 1024
IWORK_PREVIEWS = ("quicklook/preview.pdf", "preview.pdf")
# EPUB font obfuscation (not DRM): books with only these in encryption.xml are readable
FONT_OBFUSCATION = {"http://www.idpf.org/2008/embedding", "http://ns.adobe.com/pdf/enc#RC"}
EPUB_IMAGE_TYPES = {"JPEG": ("jpg", "image/jpeg"), "PNG": ("png", "image/png"), "GIF": ("gif", "image/gif")}

BLOCK_TAGS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "pre", "blockquote", "table",
              "hr", "section", "article", "header", "footer", "nav", "aside", "figure", "figcaption", "main",
              "body", "dl", "dt", "dd", "address", "center"}
DROP_TAGS = {"script", "style", "iframe", "object", "embed", "form", "input", "button", "select", "textarea",
             "link", "meta", "noscript", "svg", "math", "head", "title", "template", "canvas", "audio", "video"}
KEEP_TAGS = BLOCK_TAGS | {"a", "b", "strong", "i", "em", "u", "s", "del", "code", "kbd", "samp", "br", "img",
                          "span", "sub", "sup", "small", "mark", "thead", "tbody", "tfoot", "tr", "th", "td",
                          "caption", "q", "cite", "abbr", "time", "var", "ins"}
KEEP_ATTRS = {"href", "src", "alt", "title", "colspan", "rowspan", "start"}

PAGE_CSS = """body { font-family: -apple-system, "Segoe UI", Helvetica, Arial, sans-serif; line-height: 1.5;
  max-width: 46em; margin: 2em auto; padding: 0 1em; color: #222; }
pre, code { font-family: Menlo, Consolas, monospace; font-size: 0.92em; }
pre { background: #f5f5f5; padding: 0.8em; overflow-x: auto; white-space: pre-wrap; }
table { border-collapse: collapse; } td, th { border: 1px solid #bbb; padding: 0.3em 0.6em; }
img { max-width: 100%; height: auto; } blockquote { border-left: 3px solid #ccc; margin-left: 0; padding-left: 1em; }
"""


@dataclass
class Doc:
    title: str
    body: str                                   # HTML body fragment
    base: Path | None = None                    # folder relative picture links resolve against
    images: dict[str, bytes] = field(default_factory=dict)  # pictures that aren't files on disk


# ---------------------------------------------------------------- dispatch
def targets_for(cat: str, src_ext: str) -> list[str]:
    ext = "." + src_ext.lower().lstrip(".")
    if cat == "sheet":
        own = {".xlsx": "XLSX", ".xlsm": "XLSX", ".ods": "ODS", ".csv": "CSV", ".tsv": "TSV"}.get(ext)
        return [t for t in SHEET_TARGETS if t != own]
    if cat == "text":
        if ext in HTML_EXTS:
            return HTML_TARGETS
        return PLAIN_TARGETS if ext == ".txt" else TEXT_TARGETS
    return {"slides": SLIDES_TARGETS, "epub": EPUB_TARGETS, "markdown": MARKDOWN_TARGETS,
            "rtf": RTF_TARGETS, "pages": PAGES_TARGETS, "pdf": PDF_TARGETS, "docx": DOCX_TARGETS}.get(cat, [])


def convert(src: str | Path, cat: str, target: str) -> list[Path]:
    src = Path(src)
    target = target.upper()
    if target not in targets_for(cat, src.suffix):
        raise ValueError(f"Can't convert {cat} to {target}")
    if cat == "sheet":
        return sheet_to(src, target)
    if cat == "slides":
        from . import slides
        return slides.slides_to(src, target)
    if cat == "pages":
        return pages_to_pdf(src)
    if cat == "pdf" and target == "XLSX":
        return pdf_to_xlsx(src)
    if cat == "pdf" and target == "PPTX":
        from . import slides
        return slides.pdf_to_pptx(src)
    if cat == "epub" and target == "PDF":
        return epub_to_pdf(src)
    return [write_doc(read_doc(src, cat), src, target)]


# ---------------------------------------------------------------- spreadsheets
def sheet_to(src: Path, target: str) -> list[Path]:
    book = sheets.read_sheets(src)
    if target in ("CSV", "TSV"):
        return sheets.write_delimited(src, book, "." + target.lower())
    out = sibling(src, ext="." + target.lower())
    if target == "XLSX":
        sheets.write_xlsx(book, out)
    elif target == "ODS":
        sheets.write_ods(book, out)
    elif target == "PDF":
        body = sheets.sheets_html(book)
        atomic_save(lambda dst: html_to_pdf(body, dst), out)
    else:  # DOCX
        _write_docx(Doc(src.stem, sheets.sheets_html(book)), out)
    return [out]


def pdf_to_xlsx(src: Path) -> list[Path]:
    """Each table PyMuPDF finds becomes a sheet. A PDF without ruled tables gets one sheet
    per page with a row per line of text, so nothing is lost."""
    found: list[sheets.Sheet] = []
    lines: list[sheets.Sheet] = []
    with open_pdf(src) as doc:
        for pno, page in enumerate(doc):
            tables = page.find_tables().tables
            for ti, table in enumerate(tables):
                rows = [[c if c is not None else "" for c in row] for row in table.extract()]
                name = f"Page {pno + 1}" + (f" table {ti + 1}" if len(tables) > 1 else "")
                found.append(sheets.Sheet(name, sheets.trim(rows), from_text=True))
            text = [[ln] for ln in page.get_text("text").splitlines() if ln.strip()]
            lines.append(sheets.Sheet(f"Page {pno + 1}", text, from_text=True))
    book = [s for s in found if s.rows] or [s for s in lines if s.rows]
    if not book:
        raise UserError("This PDF has no text to put in a spreadsheet. It may be a scan.")
    out = sibling(src, ext=".xlsx")
    sheets.write_xlsx(book, out)
    return [out]


# ---------------------------------------------------------------- readers
def read_doc(src: Path, cat: str) -> Doc:
    if cat == "docx":
        return read_docx(src)
    if cat == "markdown":
        return read_markdown(src)
    if cat == "epub":
        return read_epub(src)
    if cat == "rtf":
        return read_rtf(src)
    if cat == "pdf":
        return read_pdf(src)
    if src.suffix.lower() in HTML_EXTS:
        return read_html(src)
    return read_plain(src)


def read_docx(src: Path) -> Doc:
    from .convert import DOCX_DAMAGED, docx_to_html
    images: dict[str, bytes] = {}
    try:
        body = docx_to_html(src, images)
    except UserError:
        raise
    except Exception as e:  # python-docx raises zip, XML and key errors on damaged files
        raise UserError(DOCX_DAMAGED) from e
    return Doc(_title_from(body, src), body, images=images)


def read_markdown(src: Path) -> Doc:
    from markdown_it import MarkdownIt
    # CommonMark plus GitHub's tables and ~~strikethrough~~; raw HTML is allowed (it's cleaned later)
    md = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])
    body = md.render(read_text(src))
    return Doc(_title_from(body, src), body, base=src.parent)


def read_plain(src: Path) -> Doc:
    """Plain text: paragraphs at blank lines, line breaks kept."""
    paras = re.split(r"\n\s*\n", read_text(src).replace("\r\n", "\n").replace("\r", "\n"))
    body = "\n".join("<p>" + "<br/>".join(html.escape(ln) for ln in p.strip("\n").split("\n")) + "</p>"
                     for p in paras if p.strip())
    return Doc(src.stem, body)


def read_html(src: Path) -> Doc:
    import lxml.html
    raw = src.read_bytes()
    if not raw.strip():
        raise UserError("This file is empty.")
    tree = lxml.html.document_fromstring(raw)
    title = (tree.findtext(".//title") or "").strip() or src.stem
    return Doc(title, _inner_html(tree.find("body") if tree.find("body") is not None else tree), base=src.parent)


def read_rtf(src: Path) -> Doc:
    """RTF (or a TextEdit .rtfd bundle). macOS's textutil keeps formatting; elsewhere the text
    is extracted with striprtf."""
    path = src / "TXT.rtf" if src.is_dir() else src
    if not path.is_file():
        raise UserError("This RTFD bundle has no text inside.")
    raw = path.read_bytes()
    if not raw.lstrip().startswith(b"{\\rtf"):
        raise UserError("This doesn't look like an RTF file.")
    page = None
    if sys.platform == "darwin":
        from . import platform_mac
        page = platform_mac.textutil_html(src)
    if page:
        import lxml.html
        tree = lxml.html.document_fromstring(page.encode("utf-8"))
        body = tree.find("body")
        return Doc(src.stem, _inner_html(body if body is not None else tree), base=src if src.is_dir() else None)
    from striprtf.striprtf import rtf_to_text
    text = rtf_to_text(raw.decode("latin-1"), errors="ignore")
    paras = [p for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    return Doc(src.stem, "\n".join(f"<p>{html.escape(p)}</p>" for p in paras))


def read_pdf(src: Path) -> Doc:
    """Text blocks as paragraphs (large text as headings) and embedded pictures, in page order."""
    parts, images = [], {}
    with open_pdf(src) as doc:
        title = (doc.metadata or {}).get("title") or src.stem
        for pno, page in enumerate(doc):
            for _, kind, data in pdf_page_items(page):
                if kind == "text":
                    txt, size, bold = data
                    level = heading_level(size)
                    if level:
                        parts.append(f"<h{level}>{html.escape(txt)}</h{level}>")
                    elif bold and len(txt) < 120:
                        parts.append(f"<p><b>{html.escape(txt)}</b></p>")
                    else:
                        parts.append(f"<p>{html.escape(txt)}</p>")
                else:
                    raw, _bbox, ext = data
                    name = f"page{pno + 1}-image{len(images) + 1}.{ext}"
                    images[name] = raw
                    parts.append(f'<p><img src="{name}" alt=""/></p>')
    if not parts:
        raise UserError("This PDF has no text or pictures to convert. It may be a scan.")
    return Doc(title, "\n".join(parts), images=images)


def read_epub(src: Path) -> Doc:
    import lxml.html
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    try:
        z = zipfile.ZipFile(src)
    except (zipfile.BadZipFile, OSError) as e:
        raise UserError("This EPUB is damaged and can't be opened.") from e
    with z:
        try:
            _check_epub_drm(z, parser)
            container = etree.fromstring(_zip_read(z, "META-INF/container.xml"), parser)
            rootfile = container.find(".//{urn:oasis:names:tc:opendocument:xmlns:container}rootfile")
            opf_path = rootfile.get("full-path")
            opf = etree.fromstring(_zip_read(z, opf_path), parser)
        except UserError:
            raise
        except Exception as e:  # missing parts, bad XML
            raise UserError("This EPUB is damaged and can't be opened.") from e
        ns = {"opf": "http://www.idpf.org/2007/opf", "dc": "http://purl.org/dc/elements/1.1/"}
        title = (opf.findtext(".//dc:title", namespaces=ns) or "").strip() or src.stem
        base = posixpath.dirname(opf_path)
        manifest = {}
        for item in opf.iterfind(".//opf:manifest/opf:item", ns):
            href = item.get("href")
            if item.get("id") and href:
                manifest[item.get("id")] = (posixpath.normpath(posixpath.join(base, unquote(href))),
                                            item.get("media-type", ""))
        parts, images = [], {}
        seen: set[str] = set()   # a spine can list one chapter many times over
        unreadable: set[str] = set()   # missing or oversized pictures, tried once each
        total = 0

        def counted(data: bytes) -> bytes:
            nonlocal total
            total += len(data)
            if total > MAX_BOOK_BYTES:
                raise UserError("This ebook is too large to convert.")
            return data

        for ref in opf.iterfind(".//opf:spine/opf:itemref", ns):
            path, mtype = manifest.get(ref.get("idref"), (None, ""))
            if path is None or path in seen or mtype not in ("application/xhtml+xml", "text/html"):
                continue
            seen.add(path)
            try:
                tree = lxml.html.document_fromstring(counted(_zip_read(z, path)))
            except (KeyError, etree.ParserError, ValueError):
                continue
            body = tree.find("body")
            if body is None:
                continue
            for img in list(body.iter("img")):
                ipath = posixpath.normpath(posixpath.join(posixpath.dirname(path), unquote(img.get("src", ""))))
                if ipath in images:
                    img.set("src", ipath)
                    continue
                data = None
                if ipath not in unreadable:
                    try:
                        data = _zip_read(z, ipath, MAX_IMAGE_BYTES)
                    except (KeyError, UserError):
                        unreadable.add(ipath)
                if data is None:
                    img.drop_tree()
                    continue
                images[ipath] = counted(data)
                img.set("src", ipath)
            parts.append(_inner_html(body))
    if not parts:
        raise UserError("This EPUB has no readable chapters.")
    return Doc(title, "\n".join(parts), images=images)


def _check_epub_drm(z: zipfile.ZipFile, parser):
    from lxml import etree
    if "META-INF/encryption.xml" not in z.namelist():
        return
    root = etree.fromstring(_zip_read(z, "META-INF/encryption.xml"), parser)
    for method in root.iter("{http://www.w3.org/2001/04/xmlenc#}EncryptionMethod"):
        if method.get("Algorithm") not in FONT_OBFUSCATION:
            raise UserError("This ebook is DRM-protected, so it can't be converted.")


def _zip_read(z: zipfile.ZipFile, name: str, limit: int = MAX_ZIP_ENTRY) -> bytes:
    info = z.getinfo(name)  # KeyError if missing
    if info.file_size > limit:
        raise UserError("A file inside this document is too large to read.")
    with z.open(info) as f:
        data = f.read(limit + 1)
    if len(data) > limit:  # the header lied about the size
        raise UserError("A file inside this document is too large to read.")
    return data


# ---------------------------------------------------------------- HTML helpers
def _parse(body: str):
    """The body fragment as a cleaned <div> tree: scripts, forms and unknown tags removed,
    only harmless attributes kept."""
    import lxml.html
    root = lxml.html.fragment_fromstring(CONTROL_CHARS.sub("", body) or "", create_parent="div")
    for el in list(root.iter()):
        if not isinstance(el.tag, str):  # comments and processing instructions
            if el.getparent() is not None:
                el.drop_tree()
            continue
        if el is root:
            continue
        tag = el.tag.lower()
        if tag in DROP_TAGS:
            el.drop_tree()
        elif tag not in KEEP_TAGS:
            el.drop_tag()  # keep the text, lose the element
    for el in root.iter():
        for name in list(el.attrib):
            value = el.attrib[name]
            if name == "class" and el.tag == "code" and value.startswith("language-"):
                continue  # the code block's language, kept for Markdown fences
            script_link = name == "href" and value.strip().lower().startswith(("javascript:", "vbscript:"))
            if name not in KEEP_ATTRS or script_link:
                del el.attrib[name]
    return root


def _inner_html(el) -> str:
    from lxml import etree
    text = html.escape(el.text or "")
    return text + "".join(etree.tostring(c, method="html", encoding="unicode") for c in el)


def _title_from(body: str, src: Path) -> str:
    m = re.search(r"<h1[^>]*>(.*?)</h1>", body, re.S | re.I)
    if m:
        text = html.unescape(re.sub(r"<[^>]+>", "", m.group(1))).strip()
        if text:
            return text[:200]
    return src.stem


def _image_bytes(doc: Doc, src: str) -> bytes | None:
    """The picture an <img src> points at: from the document, a data: URI, or a local file.
    Never fetched from the network."""
    if not src:
        return None
    if src in doc.images:
        return doc.images[src]
    if src.startswith("data:image/") and ";base64," in src:
        try:
            return base64.b64decode(src.split(",", 1)[1], validate=False)
        except ValueError:
            return None
    if "://" in src or src.startswith("//") or doc.base is None:
        return None
    path = _local_picture(doc.base, unquote(src.split("#")[0].split("?")[0]))
    try:
        if path is not None and path.is_file() and path.stat().st_size <= MAX_IMAGE_BYTES:
            return path.read_bytes()
    except OSError:
        pass
    return None


def _local_picture(base: Path, ref: str) -> Path | None:
    """The file a relative picture link names, only when it's inside the document's own folder.
    A downloaded page could otherwise pull any local file into the result (an absolute path,
    "../", or a symlink leading out), so those give None."""
    if not ref or PurePosixPath(ref).is_absolute() or PureWindowsPath(ref).anchor:
        return None
    try:
        root = Path(os.path.realpath(base))
        path = Path(os.path.realpath(root / ref))   # follows every symlink on the way
        path.relative_to(root)
    except (OSError, ValueError):
        return None
    return path


def _web_image(data: bytes) -> tuple[bytes, str, str] | None:
    """(bytes, extension, media type) in a format every ebook reader shows; other formats are
    converted to PNG. None if it isn't a picture."""
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data))
        fmt = im.format
        if fmt in EPUB_IMAGE_TYPES:
            ext, mime = EPUB_IMAGE_TYPES[fmt]
            return data, ext, mime
        buf = io.BytesIO()
        im.convert("RGBA").save(buf, "PNG")
        return buf.getvalue(), "png", "image/png"
    except Exception:
        return None


# ---------------------------------------------------------------- writers
def write_doc(doc: Doc, src: Path, target: str) -> Path:
    out = sibling(src, ext="." + target.lower())
    if target == "PDF":
        _write_pdf(doc, out)
    elif target == "DOCX":
        _write_docx(doc, out)
    elif target in ("MD", "TXT"):
        text = html_to_text(doc, markdown=target == "MD")
        if not text.strip():
            raise UserError("There's no text in this document to convert.")
        save_text(out, text)
    elif target == "HTML":
        save_text(out, html_page(doc))
    elif target == "EPUB":
        write_epub(doc, out)
    else:
        raise ValueError(f"Can't write {target}")
    return out


def _write_pdf(doc: Doc, out: Path) -> None:
    from lxml import etree
    root = _parse(doc.body)
    images = {}
    for i, img in enumerate(root.iter("img")):
        data = _image_bytes(doc, img.get("src", ""))
        if data is None:
            img.drop_tree()
            continue
        name = f"img{i + 1}"
        images[name] = data
        img.set("src", name)
    body = "".join(etree.tostring(c, method="html", encoding="unicode") for c in root)
    body = html.escape(root.text or "") + body
    atomic_save(lambda dst: html_to_pdf(body, dst, images), out)


def html_page(doc: Doc) -> str:
    """A standalone web page; pictures are embedded so it works anywhere it's moved."""
    from lxml import etree
    root = _parse(doc.body)
    for img in list(root.iter("img")):
        data = _image_bytes(doc, img.get("src", ""))
        pic = _web_image(data) if data else None
        if pic is None:
            if not img.get("src", "").startswith(("http://", "https://")):
                img.drop_tree()
            continue
        img.set("src", f"data:{pic[2]};base64,{base64.b64encode(pic[0]).decode('ascii')}")
    body = html.escape(root.text or "") + "".join(etree.tostring(c, method="html", encoding="unicode") for c in root)
    return ("<!DOCTYPE html>\n<html>\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
            f"<title>{html.escape(doc.title)}</title>\n<style>\n{PAGE_CSS}</style>\n</head>\n"
            f"<body>\n{body}\n</body>\n</html>\n")


# ------------------------------------------------ Markdown and plain text
class _TextWriter:
    """Turns the cleaned HTML tree into Markdown, or plain text when markdown=False."""

    def __init__(self, doc: Doc, markdown: bool):
        self.doc, self.md = doc, markdown

    def esc(self, text: str) -> str:
        text = re.sub(r"\s+", " ", text)
        if self.md:
            text = re.sub(r"([\\`*_\[\]<>])", r"\\\1", text)
        return text

    def block(self, el, depth: int = 0) -> str:
        """Markdown for an element whose children may mix inline content and blocks."""
        parts, inline = [], []

        def flush():
            text = "".join(inline).strip()
            inline.clear()
            if text:
                if self.md:
                    text = re.sub(r"^([#>+\-]|\d+[.)])", r"\\\1", text)
                parts.append(text)

        inline.append(self.esc(el.text or ""))
        for child in el:
            if child.tag in BLOCK_TAGS:
                flush()
                text = self.element(child, depth)
                if text.strip() and parts and el.tag == "li" and child.tag in ("ul", "ol"):
                    parts[-1] += "\n" + text  # a sub-list sits right under its item
                else:
                    parts.append(text)
            else:
                inline.append(self.inline(child))
            inline.append(self.esc(child.tail or ""))
        flush()
        return "\n\n".join(p for p in parts if p.strip())

    def element(self, el, depth: int) -> str:
        tag = el.tag
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            text = self.inline_content(el).strip()
            if not text:
                return ""
            return ("#" * int(tag[1]) + " " + text) if self.md else text
        if tag == "pre":
            code = el.text_content().strip("\n")
            if not self.md:
                return code
            fence = "````" if "```" in code else "```"
            lang = ""
            first = el.find("code")
            if first is not None:
                m = re.search(r"language-(\w+)", first.get("class", ""))
                lang = m.group(1) if m else ""
            return f"{fence}{lang}\n{code}\n{fence}"
        if tag in ("ul", "ol"):
            return self.list(el, depth)
        if tag == "blockquote":
            inner = self.block(el, depth)
            prefix = "> " if self.md else "    "
            return "\n".join((prefix + ln).rstrip() for ln in inner.split("\n"))
        if tag == "table":
            return self.table(el)
        if tag == "hr":
            return "---" if self.md else "----------"
        return self.block(el, depth)

    def list(self, el, depth: int) -> str:
        lines = []
        start = int(el.get("start", 1)) if (el.get("start") or "").isdigit() else 1
        n = start
        marker = "- "
        for child in el:
            if child.tag in ("ul", "ol"):  # a list nested straight inside a list: under the last item
                lines.append(_indent(self.list(child, depth + 1), " " * len(marker)))
                continue
            marker = f"{n}. " if el.tag == "ol" else "- "
            n += 1
            content = self.block(child, depth + 1) if child.tag == "li" else self.inline(child)
            content = content or ""
            first, _, rest = content.partition("\n")
            item = marker + first
            if rest:
                item += "\n" + _indent(rest, " " * len(marker))
            lines.append(item)
        return "\n".join(lines)

    def table(self, el) -> str:
        rows = []
        for tr in el.iter("tr"):
            rows.append([self.inline_content(c).strip().replace("|", "\\|" if self.md else "|")
                         for c in tr if c.tag in ("td", "th")])
        rows = [r for r in rows if r]
        if not rows:
            return ""
        if not self.md:
            return "\n".join("\t".join(r) for r in rows)
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
        lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return "\n".join(lines)

    def inline_content(self, el) -> str:
        return self.esc(el.text or "") + "".join(self.inline(c) + self.esc(c.tail or "") for c in el)

    def inline(self, el) -> str:
        tag = el.tag
        if tag == "br":
            return "  \n" if self.md else "\n"
        if tag == "img":
            alt = el.get("alt", "")
            src = el.get("src", "")
            if not self.md:
                return alt
            if not src or src in self.doc.images or src.startswith("data:"):
                return self.esc(alt)  # the picture isn't a file a Markdown link can point at
            return f"![{self.esc(alt)}]({src})"
        if tag in BLOCK_TAGS:  # a block inside inline content: keep its words
            return " " + self.inline_content(el) + " "
        inner = self.inline_content(el)
        if not self.md:
            href = el.get("href", "") if tag == "a" else ""
            return f"{inner} ({href})" if href and href.startswith(("http", "mailto:")) and href != inner else inner
        if tag in ("b", "strong"):
            return _wrap(inner, "**")
        if tag in ("i", "em", "cite"):
            return _wrap(inner, "*")
        if tag in ("s", "del"):
            return _wrap(inner, "~~")
        if tag in ("code", "kbd", "samp"):
            code = el.text_content()
            tick = "``" if "`" in code else "`"
            return f"{tick}{code}{tick}" if code.strip() else ""
        if tag == "a":
            href = el.get("href", "")
            return _wrap(inner, "", f"]({href})", "[") if href and inner.strip() else inner
        return inner


def _wrap(inner: str, mark: str, close: str | None = None, open_: str | None = None) -> str:
    """**bold** (or [link](url)) with the spaces kept outside the markers: Markdown ignores '** x**'."""
    core = inner.strip()
    if not core:
        return inner
    lead = inner[:len(inner) - len(inner.lstrip())]
    trail = inner[len(inner.rstrip()):]
    return f"{lead}{mark if open_ is None else open_}{core}{mark if close is None else close}{trail}"


def _indent(text: str, pad: str) -> str:
    return "\n".join(pad + ln if ln.strip() else ln for ln in text.split("\n"))


def html_to_text(doc: Doc, markdown: bool) -> str:
    text = _TextWriter(doc, markdown).block(_parse(doc.body))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text + "\n" if text else ""


# ------------------------------------------------ Word
class _DocxWriter:
    """Builds a Word document from the cleaned HTML tree with python-docx."""

    def __init__(self, doc: Doc):
        from docx import Document
        self.src, self.d = doc, Document()
        self.d.core_properties.title = doc.title[:255]

    def save(self, out: Path):
        atomic_save(self.d.save, out)

    def blocks(self, el, depth: int = 0, style: str | None = None):
        """Paragraphs for an element whose children may mix inline content and blocks."""
        para = None

        def current():
            nonlocal para
            if para is None:
                para = self.d.add_paragraph(style=style)
            return para

        if (el.text or "").strip():
            self.add_text(current(), el.text, {})
        for child in el:
            if child.tag in BLOCK_TAGS:
                para = None
                self.block(child, depth, style)
            else:
                if child.tag == "img" or (child.text_content() or "").strip() or child.tag == "br":
                    self.inline(current(), child, {})
            if (child.tail or "").strip():
                self.add_text(current(), child.tail, {})

    def block(self, el, depth: int, style: str | None):
        tag = el.tag
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            p = self.d.add_heading(level=int(tag[1]))
            self.inline_children(p, el, {})
        elif tag == "pre":
            from docx.shared import Pt
            p = self.d.add_paragraph()
            run = p.add_run(el.text_content().strip("\n"))
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        elif tag in ("ul", "ol"):
            self.list(el, depth)
        elif tag == "blockquote":
            self.blocks(el, depth, "Quote")
        elif tag == "table":
            self.table(el)
        elif tag == "hr":
            self.d.add_paragraph()
        elif tag == "p" or tag in ("dt", "dd", "figcaption", "address"):
            self.inline_children(self.d.add_paragraph(style=style), el, {})
        else:
            self.blocks(el, depth, style)

    def list(self, el, depth: int):
        base = "List Number" if el.tag == "ol" else "List Bullet"
        style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
        for child in el:
            if child.tag in ("ul", "ol"):
                self.list(child, depth + 1)
                continue
            p = self.d.add_paragraph(style=style)
            # the item's own words go in this paragraph; lists inside it follow as deeper items
            if (child.text or "").strip():
                self.add_text(p, child.text, {})
            for sub in child:
                if sub.tag in ("ul", "ol"):
                    self.list(sub, depth + 1)
                elif sub.tag in BLOCK_TAGS:
                    self.inline_children(p, sub, {})
                else:
                    self.inline(p, sub, {})
                if (sub.tail or "").strip():
                    self.add_text(p, sub.tail, {})

    def table(self, el):
        rows = [[c for c in tr if c.tag in ("td", "th")] for tr in el.iter("tr")]
        rows = [r for r in rows if r]
        if not rows:
            return
        width = max(len(r) for r in rows)
        t = self.d.add_table(rows=len(rows), cols=width)
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                p = t.cell(r, c).paragraphs[0]
                self.inline_children(p, cell, {"bold": True} if cell.tag == "th" else {})

    def inline_children(self, p, el, fmt: dict):
        if el.text:
            self.add_text(p, el.text, fmt)
        for child in el:
            self.inline(p, child, fmt)
            if child.tail:
                self.add_text(p, child.tail, fmt)

    def inline(self, p, el, fmt: dict):
        tag = el.tag
        if tag == "br":
            p.add_run().add_break()
            return
        if tag == "img":
            self.picture(p, el)
            return
        fmt = dict(fmt)
        if tag in ("b", "strong", "th"):
            fmt["bold"] = True
        elif tag in ("i", "em", "cite"):
            fmt["italic"] = True
        elif tag == "u":
            fmt["underline"] = True
        elif tag in ("s", "del"):
            fmt["strike"] = True
        elif tag in ("code", "kbd", "samp"):
            fmt["code"] = True
        elif tag == "a" and el.get("href"):
            fmt["link"] = el.get("href")
        self.inline_children(p, el, fmt)

    def add_text(self, p, text: str, fmt: dict):
        text = re.sub(r"\s+", " ", text)
        if not text:
            return
        if not p.runs and not p._p.xpath("./w:hyperlink"):
            text = text.lstrip()
            if not text:
                return
        if fmt.get("link"):
            _add_hyperlink(p, text, fmt["link"])
            return
        run = p.add_run(text)
        run.bold = fmt.get("bold") or None
        run.italic = fmt.get("italic") or None
        run.underline = fmt.get("underline") or None
        if fmt.get("strike"):
            run.font.strike = True
        if fmt.get("code"):
            run.font.name = "Courier New"

    def picture(self, p, el):
        from docx.shared import Inches
        data = _image_bytes(self.src, el.get("src", ""))
        pic = _web_image(data) if data else None
        if pic is None:
            return
        try:
            from PIL import Image
            w_px = Image.open(io.BytesIO(pic[0])).width
            width = Inches(min(6.0, max(0.3, w_px / 96)))
            p.add_run().add_picture(io.BytesIO(pic[0]), width=width)
        except Exception:
            pass  # a picture Word can't take shouldn't stop the document


def _add_hyperlink(p, text: str, url: str):
    """A real, clickable Word hyperlink (python-docx has no public API for adding one)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.shared import OxmlElement, qn
    if url.startswith("#"):
        p.add_run(text)
        return
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([style, color, underline])
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.extend([rpr, t])
    link.append(run)
    p._p.append(link)


def _write_docx(doc: Doc, out: Path) -> None:
    w = _DocxWriter(doc)
    w.blocks(_parse(doc.body))
    w.save(out)


# ------------------------------------------------ EPUB
def write_epub(doc: Doc, out: Path) -> None:
    """An EPUB 3 book (with an EPUB 2 table of contents for older readers). A new chapter starts
    at each top-level heading."""
    root = _parse(doc.body)
    images: dict[str, tuple[bytes, str]] = {}  # file name in the book -> (bytes, media type)
    for img in list(root.iter("img")):
        data = _image_bytes(doc, img.get("src", ""))
        pic = _web_image(data) if data else None
        if pic is None:
            img.drop_tree()
            continue
        name = f"images/image{len(images) + 1}.{pic[1]}"
        images[name] = (pic[0], pic[2])
        img.set("src", name)
    chapters = _split_chapters(root, doc.title)
    if not any(html_ for _, html_ in chapters):
        raise UserError("There's no text to put in an ebook.")
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    title = html.escape(doc.title)

    def xhtml(chapter_title: str, body: str) -> str:
        return ('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
                '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">\n'
                f'<head><meta charset="utf-8"/><title>{html.escape(chapter_title)}</title>'
                '<link rel="stylesheet" type="text/css" href="style.css"/></head>\n'
                f"<body>\n{body}\n</body>\n</html>\n")

    files = {}
    nav_items, ncx_points, manifest, spine = [], [], [], []
    for i, (ch_title, body) in enumerate(chapters, 1):
        name = f"chapter{i}.xhtml"
        files[name] = xhtml(ch_title, body).encode("utf-8")
        manifest.append(f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i}"/>')
        nav_items.append(f'<li><a href="{name}">{html.escape(ch_title)}</a></li>')
        ncx_points.append(f'<navPoint id="p{i}" playOrder="{i}"><navLabel><text>{html.escape(ch_title)}</text>'
                          f'</navLabel><content src="{name}"/></navPoint>')
    for i, (name, (_, mime)) in enumerate(images.items(), 1):
        manifest.append(f'<item id="img{i}" href="{name}" media-type="{mime}"/>')
    nav = xhtml("Contents", f'<nav epub:type="toc" id="toc"><h1>Contents</h1><ol>{"".join(nav_items)}</ol></nav>')
    opf = ('<?xml version="1.0" encoding="utf-8"?>\n'
           '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">\n'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
           f'<dc:identifier id="bookid">{book_id}</dc:identifier>\n<dc:title>{title}</dc:title>\n'
           f'<dc:language>und</dc:language>\n<meta property="dcterms:modified">{modified}</meta>\n</metadata>\n'
           '<manifest>\n<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
           '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
           '<item id="css" href="style.css" media-type="text/css"/>\n' + "\n".join(manifest) +
           '\n</manifest>\n<spine toc="ncx">\n' + "\n".join(spine) + "\n</spine>\n</package>\n")
    ncx = ('<?xml version="1.0" encoding="utf-8"?>\n'
           '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
           f'<head><meta name="dtb:uid" content="{book_id}"/></head>\n'
           f'<docTitle><text>{title}</text></docTitle>\n<navMap>{"".join(ncx_points)}</navMap>\n</ncx>\n')
    container = ('<?xml version="1.0" encoding="utf-8"?>\n'
                 '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 '</rootfiles>\n</container>\n')

    def write(dst):
        with zipfile.ZipFile(dst, "w") as z:
            # the mimetype must come first and be stored uncompressed
            z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
            z.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/toc.ncx", ncx, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/style.css", "body { line-height: 1.5; } img { max-width: 100%; } "
                       "pre { white-space: pre-wrap; } td, th { border: 1px solid #999; padding: 0.2em 0.4em; } "
                       "table { border-collapse: collapse; }", compress_type=zipfile.ZIP_DEFLATED)
            for name, data in files.items():
                z.writestr(f"OEBPS/{name}", data, compress_type=zipfile.ZIP_DEFLATED)
            for name, (data, _) in images.items():
                z.writestr(f"OEBPS/{name}", data, compress_type=zipfile.ZIP_DEFLATED)
    atomic_save(write, out)


def _split_chapters(root, title: str) -> list[tuple[str, str]]:
    """[(chapter title, XHTML body)], split at h1 headings (or h2 if the text has no h1)."""
    from lxml import etree
    children = list(root)
    level = next((h for h in ("h1", "h2") if any(c.tag == h for c in children)), None)
    chapters: list[tuple[str, list]] = [(title, [])]
    lead = html.escape(root.text or "").strip()
    if lead:
        chapters[0][1].append(f"<p>{lead}</p>")
    for c in children:
        if level and c.tag == level and (chapters[-1][1] or len(chapters) > 1):
            chapters.append((c.text_content().strip() or title, []))
        elif level and c.tag == level and not chapters[-1][1]:
            chapters[-1] = (c.text_content().strip() or title, [])
        chapters[-1][1].append(etree.tostring(c, method="xml", encoding="unicode"))
    return [(t, "\n".join(parts)) for t, parts in chapters if parts]


def epub_to_pdf(src: Path) -> list[Path]:
    """MuPDF lays the book out on Letter pages itself, keeping its styles and pictures."""
    read_epub(src)  # same checks (damaged, DRM) and messages as the other EPUB conversions
    out = sibling(src, ext=".pdf")
    try:
        with fitz.open(src) as book:
            book.layout(rect=fitz.paper_rect("letter"), fontsize=11)
            data = book.convert_to_pdf()
    except Exception as e:
        raise UserError("This EPUB couldn't be laid out as a PDF.") from e
    save_bytes(out, data)
    return [out]


# ---------------------------------------------------------------- Pages and Keynote
def iwork_to_pdf(app: str, src: Path, out: Path) -> bool:
    """Export through Keynote or Pages on a Mac that has it. False if not possible, so the
    caller can fall back to the built-in conversion."""
    if sys.platform != "darwin":
        return False
    from . import platform_mac
    tmp = out.with_name(f".~{out.stem}.export.pdf")
    try:
        if not platform_mac.iwork_export_pdf(app, src, tmp):
            return False
        with fitz.open(tmp) as check:  # make sure it's a real PDF before keeping it
            if check.page_count == 0:
                return False
        os.replace(tmp, out)
        return True
    except Exception:
        return False
    finally:
        try:
            tmp.unlink()
        except OSError:
            pass


def pages_to_pdf(src: Path) -> list[Path]:
    """A Pages document as PDF: exported by Pages when it's installed, otherwise the preview
    Pages saves inside every document (a full PDF in older files, a first-page picture in newer)."""
    out = sibling(src, ext=".pdf")
    if iwork_to_pdf("Pages", src, out):
        return [out]
    kind, data = _pages_preview(src)
    if kind == "pdf":
        try:
            with fitz.open(stream=data, filetype="pdf") as check:
                if check.page_count == 0:
                    raise ValueError("empty")
        except Exception as e:
            raise UserError("The preview inside this Pages file is damaged.") from e
        save_bytes(out, data)
    else:
        from PIL import Image
        try:
            im = Image.open(io.BytesIO(data))
            im.load()
        except Exception as e:
            raise UserError("The preview inside this Pages file is damaged.") from e
        save_image(im, out, "PDF")
    return [out]


def _pages_preview(src: Path) -> tuple[str, bytes]:
    """('pdf' or 'image', bytes) of the best preview inside a .pages file or bundle folder."""
    no_preview = ("This Pages file has no preview inside it to convert. Open it in Pages and use "
                  "File, Export To, PDF (or Word), then convert that file.")
    entries: dict[str, bytes | None] = {}
    if src.is_dir():
        for p in src.rglob("*"):
            if p.is_file():
                entries[p.relative_to(src).as_posix()] = None
        read = lambda name: (src / name).read_bytes()  # noqa: E731
        z = None
    else:
        try:
            z = zipfile.ZipFile(src)
        except (zipfile.BadZipFile, OSError) as e:
            raise UserError("This Pages file is damaged and can't be opened.") from e
        entries = {n: None for n in z.namelist()}
        read = lambda name: _zip_read(z, name)  # noqa: E731
    try:
        lower = {n.lower(): n for n in entries}
        for name in IWORK_PREVIEWS:
            if name in lower:
                return "pdf", read(lower[name])
        best, best_px = None, 0
        for low, name in lower.items():
            base = posixpath.basename(low)
            if not (base.startswith(("preview", "thumbnail")) and base.endswith((".jpg", ".jpeg", ".png"))):
                continue
            if low.count("/") > 1:
                continue
            try:
                from PIL import Image
                data = read(name)
                with Image.open(io.BytesIO(data)) as im:
                    px = im.width * im.height
            except Exception:
                continue
            if px > best_px:
                best, best_px = data, px
        if best is None:
            raise UserError(no_preview)
        return "image", best
    finally:
        if z is not None:
            z.close()


# ---------------------------------------------------------------- facts for the result card
def facts(path: str | Path) -> list[str]:
    """['3 sheets', '120 rows'] for a spreadsheet, ['12 slides'] for a deck; [] otherwise."""
    path = Path(path)
    try:
        ext = path.suffix.lower()
        if ext in (".xlsx", ".ods", ".csv", ".tsv"):
            book = [s for s in sheets.read_sheets(path) if s.rows]
            bits = [plural(len(book), "sheet")] if len(book) > 1 else []
            n = sheets.row_count(book)
            return bits + [f"{n:,} row{'s' if n != 1 else ''}"]
        if ext == ".pptx":
            from . import slides
            n = slides.slide_count(path)
            return [plural(n, "slide")]
    except Exception:
        return []
    return []


def unit_count(src: str | Path) -> tuple[int, str] | None:
    """How many sheets or slides a source has, for titles like '3 sheets → 3 CSVs'."""
    src = Path(src)
    try:
        ext = src.suffix.lower()
        if ext in (".xlsx", ".xlsm", ".xls", ".ods"):
            return len([s for s in sheets.read_sheets(src) if s.rows]), "sheets"
        if ext == ".pptx":
            from . import slides
            return slides.slide_count(src), "slides"
    except Exception:
        return None
    return None

