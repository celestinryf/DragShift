"""Format conversions. Every function takes a source path and returns a list of
output paths written next to the source."""
from __future__ import annotations

import html
import io
import os
import secrets
import sys
from pathlib import Path

from PIL import Image, ImageOps

from .util import HTML_EXTS, UserError, encrypted_ole, readable_name_errors, sibling

try:  # HEIC / HEIF support
    import pillow_heif
    pillow_heif.register_heif_opener()
    HEIF_OK = True
except Exception:  # pragma: no cover
    HEIF_OK = False

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore

IMAGE_TARGETS = ["PNG", "JPG", "WEBP", "HEIC", "TIFF", "AVIF", "BMP", "PDF", "ICO", "GIF"]
PDF_TARGETS = ["JPG", "PNG", "DOCX", "TXT", "XLSX", "PPTX", "EPUB", "MD"]
DOCX_TARGETS = ["PDF", "TXT", "MD", "HTML", "EPUB"]
TEXT_TARGETS = ["PDF", "DOCX"]
# handled by documents.py (spreadsheets, slides, ebooks, Markdown, RTF, Pages)
DOCUMENT_CATEGORIES = ("sheet", "slides", "epub", "markdown", "rtf", "pages")

PIL_FORMAT = {"PNG": "PNG", "JPG": "JPEG", "WEBP": "WEBP", "HEIC": "HEIF", "TIFF": "TIFF",
              "AVIF": "AVIF", "BMP": "BMP", "PDF": "PDF", "ICO": "ICO", "GIF": "GIF"}
NO_ALPHA = {"JPG", "PDF", "BMP"}

DOCX_DAMAGED = "This Word file is damaged or isn't really a .docx file."
DOCX_LOCKED = "This Word file is password-protected. Remove the password first, then try again."
# Word's list formats (w:numFmt) that count; anything else ("bullet", "none", a picture) is a bullet
ORDERED_FORMATS = {"decimal", "decimalZero", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman",
                   "ordinal", "cardinalText", "ordinalText", "decimalEnclosedCircle", "decimalEnclosedParen",
                   "decimalEnclosedFullstop", "decimalFullWidth", "decimalHalfWidth"}


# ---------------------------------------------------------------- images
def open_image(path: str | Path) -> Image.Image:
    im = Image.open(path)
    im.seek(0)
    im = ImageOps.exif_transpose(im)  # honour camera rotation
    if im.mode in ("P", "LA", "PA"):
        im = im.convert("RGBA")
    elif im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    return im


def open_preview(path: str | Path, max_side: int = 1400) -> Image.Image:
    """Fast, reduced-size decode for on-screen previews (JPEG decodes at 1/2, 1/4, 1/8 scale)."""
    im = Image.open(path)
    im.seek(0)
    if im.format == "JPEG":
        im.draft("RGB", (max_side, max_side))
    im = ImageOps.exif_transpose(im)
    if im.mode in ("P", "LA", "PA"):
        im = im.convert("RGBA")
    elif im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    if max(im.size) > max_side:
        im.thumbnail((max_side, max_side), Image.LANCZOS)
    return im


def open_pdf(path: str | Path):
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise UserError("This PDF is password-protected. Remove the password first.")
    return doc


def flatten(im: Image.Image, bg=(255, 255, 255)) -> Image.Image:
    if im.mode == "RGBA":
        base = Image.new("RGB", im.size, bg)
        base.paste(im, mask=im.split()[3])
        return base
    return im.convert("RGB") if im.mode != "RGB" else im


def save_image(im: Image.Image, out: Path, target: str, quality: int | None = None, exif=None) -> None:
    from .util import CONFIG
    quality = quality or int(CONFIG.get("jpg_quality", 90))
    target = target.upper()
    fmt = PIL_FORMAT[target]
    if target in NO_ALPHA:
        im = flatten(im)
    kw: dict = {}
    if target in ("JPG", "WEBP", "AVIF", "HEIC"):
        kw["quality"] = quality
    if target == "JPG":
        kw.update(optimize=True, progressive=True)
    if target == "PNG":
        kw["optimize"] = True
    if target == "ICO":
        side = min(256, max(im.size))
        im = ImageOps.contain(im.convert("RGBA"), (side, side))
        kw["sizes"] = [(s, s) for s in (16, 24, 32, 48, 64, 128, 256) if s <= side]
    if target == "PDF":
        kw["resolution"] = 150.0
    if exif and target in ("JPG", "WEBP", "PNG", "TIFF"):
        kw["exif"] = exif
    atomic_save(lambda dst: im.save(dst, fmt, **kw), out)


def temp_path(out: str | Path) -> Path:
    """A hidden, unused temp name in `out`'s folder. Short and random rather than built from the
    result's name, so a long name never pushes it past the file system's 255-byte limit."""
    folder = Path(out).parent
    while True:
        tmp = folder / f".~{secrets.token_hex(6)}.part"
        if not tmp.exists() and not tmp.is_symlink():
            return tmp


def atomic_save(write, out, tmp: Path | None = None):
    """Write to a hidden temp file (`tmp`, or a fresh temp_path), then rename: Explorer never
    shows a half-written file."""
    if not isinstance(out, (str, Path)):
        write(out)  # in-memory buffer
        return
    out = Path(out)
    tmp = tmp or temp_path(out)
    try:
        with readable_name_errors():
            write(tmp)
            os.replace(tmp, out)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


def save_bytes(path: str | Path, data: bytes) -> None:
    atomic_save(lambda dst: Path(dst).write_bytes(data), path)


def save_text(path: str | Path, text: str, newline: str | None = None) -> None:
    """UTF-8 text; `newline` as for open() (None writes the platform's line ends)."""
    atomic_save(lambda dst: Path(dst).write_text(text, encoding="utf-8", newline=newline), path)


def image_to(src: str | Path, target: str) -> list[Path]:
    target = target.upper()
    im = open_image(src)
    out = sibling(src, ext="." + target.lower())
    save_image(im, out, target)
    return [out]


def images_to_single_pdf(srcs: list[str | Path]) -> list[Path]:
    pages = [flatten(open_image(s)) for s in srcs]
    out = sibling(srcs[0], "Combined", ".pdf")
    pages[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=pages[1:])
    return [out]


# ---------------------------------------------------------------- PDF
def pdf_to_images(src: str | Path, target: str, dpi: int | None = None) -> list[Path]:
    from .util import CONFIG
    dpi = dpi or int(CONFIG.get("pdf_dpi", 200))
    target = target.upper()
    outs = []
    with open_pdf(src) as doc:
        n = doc.page_count
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            label = f"page {i + 1}" if n > 1 else ""
            out = sibling(src, label, "." + target.lower())
            save_image(im, out, target)
            outs.append(out)
    return outs


def pdf_to_txt(src: str | Path) -> list[Path]:
    with open_pdf(src) as doc:
        text = "\n\f\n".join(page.get_text("text") for page in doc)
    out = sibling(src, ext=".txt")
    save_text(out, text)
    return [out]


def heading_level(size: float) -> int | None:
    """Large PDF text reads as a heading: 20 pt and up is level 1, 16 pt and up level 2."""
    if size >= 20:
        return 1
    if size >= 16:
        return 2
    return None


def pdf_page_items(page) -> list[tuple]:
    """A page's text blocks and images in reading order (top to bottom):
    (y, "text", (text, size, bold)) or (y, "img", (raw bytes, bbox, extension))."""
    items = []
    for b in page.get_text("dict")["blocks"]:
        y = b["bbox"][1]
        if b.get("type") == 0:
            lines = []
            size = 11.0
            bold = False
            for ln in b["lines"]:
                spans = ln["spans"]
                if spans:
                    size = max(size, max(s["size"] for s in spans)) if not lines else size
                    bold = bold or any("Bold" in s["font"] for s in spans)
                lines.append("".join(s["text"] for s in spans))
            txt = " ".join(l.strip() for l in lines).strip()
            if txt:
                items.append((y, "text", (txt, size, bold)))
        elif b.get("type") == 1 and b.get("image"):
            items.append((y, "img", (b["image"], b["bbox"], b.get("ext") or "png")))
    items.sort(key=lambda t: t[0])
    return items


def pdf_to_docx(src: str | Path) -> list[Path]:
    """Text blocks become paragraphs, embedded images are kept, one Word page per PDF page."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches

    d = Document()
    with open_pdf(src) as doc:
        for pno, page in enumerate(doc):
            if pno:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            for _, kind, data in pdf_page_items(page):
                if kind == "text":
                    txt, size, bold = data
                    level = heading_level(size)
                    if level:
                        d.add_heading(txt, level=level)
                    else:
                        p = d.add_paragraph()
                        r = p.add_run(txt)
                        r.bold = bold and len(txt) < 120
                else:
                    raw, bbox, _ext = data
                    width_in = min(6.5, max(0.5, (bbox[2] - bbox[0]) / 72))
                    try:
                        d.add_picture(io.BytesIO(raw), width=Inches(width_in))
                    except Exception:
                        pass
    out = sibling(src, ext=".docx")
    d.save(out)
    return [out]


# ---------------------------------------------------------------- DOCX
def check_docx(src: str | Path) -> None:
    """UserError for a password-protected Word file (Office stores those encrypted, not as a zip)."""
    if encrypted_ole(src):
        raise UserError(DOCX_LOCKED)


def open_docx(src: str | Path):
    """The python-docx Document, or a UserError with a plain sentence for a locked or damaged file."""
    from docx import Document
    check_docx(src)
    try:
        return Document(src)
    except Exception as e:  # python-docx raises zip, XML and key errors on damaged files
        raise UserError(DOCX_DAMAGED) from e


def _docx_run_html(r) -> str:
    t = html.escape(r.text)
    if not t:
        return ""
    if r.bold:
        t = f"<b>{t}</b>"
    if r.italic:
        t = f"<i>{t}</i>"
    if r.underline:
        t = f"<u>{t}</u>"
    return t


def _docx_images_html(r, part, images: dict) -> str:
    """<img> tags for the pictures inside run `r`; their bytes go into `images` by name."""
    tags = []
    for rid in r._r.xpath(".//a:blip/@r:embed"):
        target = part.related_parts.get(rid)
        if target is None:
            continue
        name = f"docx-image-{len(images) + 1}{Path(str(target.partname)).suffix}"
        images[name] = target.blob
        tags.append(f'<img src="{name}"/>')
    return "".join(tags)


def _docx_inline_html(p, images: dict | None = None) -> str:
    """A paragraph's runs and links as inline HTML. Pictures are kept only when `images` is
    given (it collects their bytes)."""
    from docx.text.hyperlink import Hyperlink

    parts = []
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            inner = "".join(_docx_run_html(r) for r in item.runs)
            url = item.address
            parts.append(f'<a href="{html.escape(url)}">{inner}</a>' if url and inner else inner)
            continue
        parts.append(_docx_run_html(item))
        if images is not None:
            parts.append(_docx_images_html(item, p.part, images))
    return "".join(parts)


def _docx_paragraph_html(p, images: dict | None = None) -> str:
    body = _docx_inline_html(p, images) or "&nbsp;"
    style = (p.style.name or "").lower() if p.style is not None else ""
    if style.startswith("heading"):
        lvl = "".join(c for c in style if c.isdigit()) or "1"
        lvl = str(min(int(lvl), 6))
        return f"<h{lvl}>{body}</h{lvl}>"
    if style == "title":
        return f"<h1>{body}</h1>"
    return f"<p>{body}</p>"


def _docx_list_info(p) -> tuple[str, int] | None:
    """('ul' or 'ol', nesting level) for a list paragraph, None for anything else."""
    style = (p.style.name or "").lower() if p.style is not None else ""
    ppr = p._p.pPr
    num = ppr.numPr if ppr is not None else None
    numbered = num is not None and num.numId is not None and num.numId.val != 0  # 0 = numbering off
    if "list" not in style and not numbered:
        return None
    if numbered and num.ilvl is not None and num.ilvl.val is not None:
        level = int(num.ilvl.val)
    else:  # "List Bullet 2" -> level 1
        digits = "".join(c for c in style if c.isdigit())
        level = int(digits) - 1 if digits else 0
    fmt = _docx_num_format(p, num.numId.val, level) if numbered else None
    ordered = fmt in ORDERED_FORMATS if fmt is not None else "number" in style
    return ("ol" if ordered else "ul"), max(0, min(level, 8))


def _docx_num_format(p, num_id: int, level: int) -> str | None:
    """The w:numFmt ('decimal', 'bullet'...) Word uses for list `num_id` at `level`, from the
    numbering part; None when the file doesn't say."""
    try:
        numbering = p.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):  # no numbering part
        return None
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    num = next((n for n in numbering.iterchildren(f"{w}num") if n.get(f"{w}numId") == str(num_id)), None)
    if num is None:
        return None
    # a list can override single levels of the shared definition it points to
    for override in num.iterchildren(f"{w}lvlOverride"):
        if override.get(f"{w}ilvl") == str(level):
            fmt = override.find(f"{w}lvl/{w}numFmt")
            if fmt is not None:
                return fmt.get(f"{w}val")
    abstract = num.find(f"{w}abstractNumId")
    if abstract is None:
        return None
    for definition in numbering.iterchildren(f"{w}abstractNum"):
        if definition.get(f"{w}abstractNumId") == abstract.get(f"{w}val"):
            for lvl in definition.iterchildren(f"{w}lvl"):
                if lvl.get(f"{w}ilvl") == str(level):
                    fmt = lvl.find(f"{w}numFmt")
                    return fmt.get(f"{w}val") if fmt is not None else None
    return None


def docx_to_html(src: str | Path, images: dict | None = None) -> str:
    """Body HTML for a Word file. Pass a dict as `images` to keep pictures (filled by name)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = open_docx(src)
    out = []
    lists: list[str] = []  # open <ul>/<ol> tags, one per nesting level

    def close_lists(depth: int = 0):
        while len(lists) > depth:
            out.append(f"</{lists.pop()}>")

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            info = _docx_list_info(para)
            if info is None:
                close_lists()
                out.append(_docx_paragraph_html(para, images))
                continue
            kind, level = info
            close_lists(level + 1)
            if len(lists) == level + 1 and lists[-1] != kind:
                close_lists(level)
            while len(lists) < level + 1:
                out.append(f"<{kind}>")
                lists.append(kind)
            out.append(f"<li>{_docx_inline_html(para, images) or '&nbsp;'}</li>")
        elif tag == "tbl":
            close_lists()
            t = Table(child, d)
            rows = []
            for row in t.rows:
                cells = "".join(f"<td>{html.escape(c.text)}</td>" for c in row.cells)
                rows.append(f"<tr>{cells}</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    close_lists()
    return "\n".join(out)


def docx_to_txt(src: str | Path) -> list[Path]:
    d = open_docx(src)
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append("\t".join(c.text for c in row.cells))
    out = sibling(src, ext=".txt")
    save_text(out, "\n".join(lines))
    return [out]


def _word_to_pdf(src: Path, out: Path) -> bool:
    """Use Microsoft Word (if installed) for a perfect-fidelity export."""
    if sys.platform != "win32":
        return False
    try:
        import win32com.client  # type: ignore
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        try:
            doc = word.Documents.Open(str(src.resolve()), ReadOnly=True)
            doc.SaveAs(str(out.resolve()), FileFormat=17)  # wdFormatPDF
            doc.Close(False)
        finally:
            word.Quit()
        return out.exists()
    except Exception:
        return False


CSS = """
* { font-family: sans-serif; }
body { font-size: 11pt; line-height: 1.35; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
p { margin: 0 0 6pt 0; }
pre { font-family: monospace; font-size: 9.5pt; white-space: pre-wrap; }
table { border-collapse: collapse; margin: 6pt 0; }
td, th { border: 1px solid #888; padding: 3pt 5pt; font-size: 10pt; }
th { font-weight: bold; background-color: #eeeeee; text-align: left; }
"""


def html_to_pdf(body_html: str, out: Path, images: dict | None = None) -> None:
    """Lay out HTML on Letter pages. Pictures come only from `images` (name -> bytes)."""
    archive = None
    if images:
        archive = fitz.Archive()
        for name, data in images.items():
            archive.add((data, name))
    story = fitz.Story(html=body_html, user_css=CSS, archive=archive)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def docx_to_pdf(src: str | Path) -> list[Path]:
    src = Path(src)
    check_docx(src)  # before Word, which would stop and ask for the password
    out = sibling(src, ext=".pdf")
    if not _word_to_pdf(src, out):
        images: dict = {}
        body = docx_to_html(src, images)
        atomic_save(lambda dst: html_to_pdf(body, dst, images), out)
    return [out]


# ---------------------------------------------------------------- text
def read_text(src: str | Path) -> str:
    raw = Path(src).read_bytes()
    # UTF-16 only with its byte-order mark: without one, almost any even-length Windows-1252
    # file "decodes" as UTF-16 and comes out as Chinese-looking garbage
    utf16 = ("utf-16",) if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ()
    for enc in ("utf-8-sig", *utf16, "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def text_to_pdf(src: str | Path) -> list[Path]:
    out = sibling(src, ext=".pdf")
    html_to_pdf(f"<pre>{html.escape(read_text(src))}</pre>", out)
    return [out]


def text_to_docx(src: str | Path) -> list[Path]:
    from docx import Document
    d = Document()
    for line in read_text(src).splitlines():
        d.add_paragraph(line)
    out = sibling(src, ext=".docx")
    d.save(out)
    return [out]


# ---------------------------------------------------------------- dispatch
def targets_for(cat: str, src_ext: str) -> list[str]:
    src_ext = src_ext.lower().lstrip(".")
    same = {"jpeg": "JPG", "jpg": "JPG", "tif": "TIFF", "tiff": "TIFF", "heif": "HEIC"}
    cur = same.get(src_ext, src_ext.upper())
    if cat == "image":
        t = [x for x in IMAGE_TARGETS if x != cur]
        if not HEIF_OK:
            t = [x for x in t if x != "HEIC"]
        return t
    if cat in ("video", "audio"):
        from . import media
        return media.targets_for(cat, src_ext)
    if cat == "archive":
        from . import archives
        return archives.targets_for(src_ext)
    if cat == "subtitle":
        from . import subtitles
        return subtitles.targets_for(src_ext)
    if cat == "svg":
        from . import svg
        return svg.targets_for()
    if cat == "text" or cat in DOCUMENT_CATEGORIES:
        from . import documents
        return documents.targets_for(cat, src_ext)
    return {"pdf": PDF_TARGETS, "docx": DOCX_TARGETS}.get(cat, [])


def convert(src: str | Path, cat: str, target: str) -> list[Path]:
    target = target.upper()
    if cat == "image":
        return image_to(src, target)
    if cat == "pdf":
        if target in ("JPG", "PNG"):
            return pdf_to_images(src, target)
        if target == "TXT":
            return pdf_to_txt(src)
        if target == "DOCX":
            return pdf_to_docx(src)
    if cat == "docx" and target in ("PDF", "TXT"):
        return docx_to_pdf(src) if target == "PDF" else docx_to_txt(src)
    if cat == "text" and target in ("PDF", "DOCX") and Path(src).suffix.lower() not in HTML_EXTS:
        return text_to_pdf(src) if target == "PDF" else text_to_docx(src)
    if cat in ("video", "audio"):
        from . import media  # imported here: media imports atomic_save from this module
        return [media.convert(src, target)]
    if cat == "archive":    # these modules import helpers from this one too
        from . import archives
        return [archives.convert(src, target)]
    if cat == "subtitle":
        from . import subtitles
        return subtitles.convert(src, target)
    if cat == "svg":
        from . import svg
        return svg.convert(src, target)
    if cat in ("pdf", "docx", "text") or cat in DOCUMENT_CATEGORIES:
        from . import documents  # imported here: documents imports helpers from this module
        return documents.convert(src, cat, target)
    raise ValueError(f"Can't convert {cat} to {target}")
